"""Maths duplicate detection on the 7 typologies never tested in Maths.

Maths has only ever been probed with True/False, MCQ, Fill in the Blank, VSAQ
and Short Answer. English covered the other seven, but duplicate behaviour
differs sharply between the two subjects (94% caught in English, 50% in Maths),
so English coverage does not transfer.

  Sheet JA - baseline, 22 rows -> 14 items: Match the Following, Assertion and
             Reasoning, FA Activity, Free Response, Long Answer, Case Based and
             Source Based, two of each.
  Sheet JB - a paraphrase of every JA question. Everything else is held
             identical, so all 14 items are the same question reworded and all
             14 should be flagged as duplicates.

Chapter CH-13 (Data Handling) is fresh - no earlier sheet used it - and it
suits sorting, matching, activity and case-based content naturally.
"""
import json
from pathlib import Path

from openpyxl import load_workbook

from tools.name_swap_sheet_builder import build_sheet, row

OUTPUT_DIR = Path("data/qar_semantic")
TEMPLATE = Path("data/upload_templates/sme_sheet.xlsx")

CURRICULUM = {
    "grade": "Grade 1", "subject": "Mathematics", "book": "Book 1", "unit": None,
    "chapter": "CH-13: So Many Toys (Data Handling)",
}
BLOOM_EXP = "The learner sorts, counts and compares groups of everyday objects."


def _curriculum():
    wb = load_workbook(TEMPLATE)
    ws = wb["Items"]
    comps = [ws.cell(row=r, column=111).value for r in range(2, 17)]
    los = {}
    for i, c in enumerate(comps):
        if c:
            los[c] = [ws.cell(row=rr, column=114 + i).value
                      for rr in range(2, ws.max_row + 1)
                      if ws.cell(row=rr, column=114 + i).value]
    wb.close()
    sort = next(c for c in comps if c and c.startswith("Sorts objects into groups"))
    prob = next(c for c in comps if c and c.startswith("Formulates and solves"))
    num = next(c for c in comps if c and c.startswith("Recognises and uses numerals"))
    cat = next(c for c in comps if c and c.startswith("Observes and understands different"))
    return ((sort, los[sort][0]), (prob, los[prob][0]), (prob, los[prob][1]),
            (num, los[num][0]), (cat, los[cat][0]))


SORT, PROB_A, PROB_B, NUM, CAT = _curriculum()

# sno, typology, (comp, lo), blooms, marks, question, answer, explanation, options
JA = [
    ("1", "Match the Following", SORT, "Understanding", "2",
     "Nikhil is sorting his toys. Match each toy to the feature it has.",
     "1-A, 2-B, 3-C", "A car has wheels, a ball rolls and bounces, a doll has arms and legs.",
     ["Car|Ball|Doll", "Has wheels|Rolls and bounces|Has arms and legs"]),
    ("2", "Match the Following", CAT, "Understanding", "2",
     "Priya is grouping her toys by colour. Match each toy to its colour.",
     "1-A, 2-B, 3-C", "The kite is red, the drum is blue and the top is green.",
     ["Red kite|Blue drum|Green top", "Red|Blue|Green"]),
    ("3", "Assertion and Reasoning", SORT, "Analysing", "1",
     "Assertion (A): Nikhil put the car and the bus in the same box. Reason (R): Both toys have wheels.",
     "A", "Both statements are true and having wheels is exactly why they group together.", None),
    ("4", "Assertion and Reasoning", SORT, "Analysing", "1",
     "Assertion (A): Priya says a ball belongs with the toys that have wheels. Reason (R): Toys can be sorted by their shape.",
     "D", "A ball has no wheels so A is false, while R is a true statement.", None),
    ("5", "FA Activity", PROB_A, "Applying", "2",
     "Nikhil wants to count his toys. Sort your toys into two groups and count how many are in each group.",
     "The toys are sorted into two groups and the count of each group is given.",
     "The activity is complete when both groups are counted.", None),
    ("6", "FA Activity", PROB_A, "Applying", "2",
     "Priya is making a toy chart. Draw a block for every toy you own and show which kind you have most of.",
     "A block chart is drawn and the most common kind of toy is named.",
     "The activity is complete when the chart is drawn and the largest group named.", None),
    ("7", "Free Response", PROB_B, "Creating", "3",
     "Nikhil asks how you would find out how many toys are in your cupboard. Write two or three sentences.",
     "I would take the toys out, put them in a line and count them one by one.",
     "Any workable counting strategy is acceptable.", None),
    ("8", "Free Response", PROB_B, "Evaluating", "3",
     "Priya wants to know which toy you have the most of. Write a few sentences about how you would check.",
     "I would sort the toys into groups by kind and see which group has the most.",
     "Any answer that involves grouping and comparing is acceptable.", None),
    ("9", "Long Answer Question", SORT, "Understanding", "4",
     "Nikhil has to explain his toy sorting to the class. Describe how you would sort a box of toys into groups and what rule you would use.",
     "I would pick one rule, such as colour or what the toy is made of, then put every toy that follows the rule into the same group and count each group.",
     "The answer should name a sorting rule and describe applying it.", None),
    ("10", "Long Answer Question", NUM, "Understanding", "4",
     "Priya is writing about her toy count. Describe how you would compare two groups of toys to say which has more.",
     "I would count each group, then compare the two numbers; the group with the bigger number has more toys.",
     "The answer should involve counting both groups and comparing the totals.", None),

    ("11", "Case Based Question", PROB_A, "Understanding", "3",
     "Nikhil emptied his toy box and found 6 cars, 4 balls and 3 dolls.",
     None, "The case sets up two questions about Nikhil's toys.", None),
    ("11.1", "Multiple Choice Question", PROB_A, "Applying", "1",
     "How many toys did Nikhil find altogether?", "13",
     "6 cars, 4 balls and 3 dolls make 13 toys.", ["11", "12", "13", "14"]),
    ("11.2", "Short Answer Question", PROB_A, "Analysing", "2",
     "Which kind of toy did Nikhil have the most of, and how many?", "Cars, and he had 6 of them.",
     "6 is the largest of the three counts.", None),

    ("12", "Case Based Question", NUM, "Understanding", "3",
     "Priya sorted her toys and found 8 blue toys and 5 red toys.",
     None, "The case sets up two questions about Priya's toys.", None),
    ("12.1", "Very Short Answer Question", NUM, "Applying", "1",
     "How many more blue toys than red toys does Priya have?", "3",
     "8 take away 5 leaves 3.", None),
    ("12.2", "Short Answer Question", NUM, "Understanding", "2",
     "Explain how you worked out the difference between the two groups.",
     "Take 5 away from 8, which leaves 3.",
     "Subtracting the smaller count from the larger gives the difference.", None),

    ("13", "Source Based Question", PROB_A, "Understanding", "3",
     "Read the table: Nikhil's class counted their toys. Cars 7, Balls 5, Kites 2.",
     None, "The table supplies the counts for both questions.", None),
    ("13.1", "Multiple Choice Question", PROB_A, "Remembering", "1",
     "According to the table, which toy was counted most?", "Cars",
     "Cars were counted 7 times, more than balls or kites.", ["Cars", "Balls", "Kites", "None of them"]),
    ("13.2", "Short Answer Question", PROB_A, "Applying", "2",
     "According to the table, how many toys were counted in total?", "14",
     "7 cars, 5 balls and 2 kites make 14 toys.", None),

    ("14", "Source Based Question", SORT, "Analysing", "3",
     "Read the note: Priya keeps her soft toys on the shelf and her wooden toys in the box.",
     None, "The note describes how Priya has sorted her toys.", None),
    ("14.1", "Very Short Answer Question", SORT, "Remembering", "1",
     "According to the note, where does Priya keep her wooden toys?", "In the box",
     "The note says the wooden toys go in the box.", None),
    ("14.2", "Short Answer Question", SORT, "Analysing", "2",
     "According to the note, what rule has Priya used to sort her toys?",
     "She sorted them by what the toys are made of.",
     "Soft versus wooden is a material rule.", None),
]

PARA = {
    "1": "Nikhil is sorting his toys. Link every toy with the feature that it has.",
    "2": "Priya is putting her toys into colour groups. Link each toy with the colour it is.",
    "3": "Assertion (A): The car and the bus were placed in one box by Nikhil. Reason (R): Wheels are found on both of those toys.",
    "4": "Assertion (A): Priya claims a ball should sit with the wheeled toys. Reason (R): Shape is one way of sorting toys.",
    "5": "Nikhil would like to count what he owns. Put your toys into two groups and work out how many are in each.",
    "6": "Priya is preparing a toy poster. Put one block down for each toy you have and show which type you own the most of.",
    "7": "Nikhil asks how you would work out the number of toys sitting in your cupboard. Write two or three sentences.",
    "8": "Priya would like to know which type of toy you own most of. Write a few sentences explaining how you would find out.",
    "9": "Nikhil must tell the class how he sorted his toys. Write about the way you would split a box of toys into groups and the rule you would follow.",
    "10": "Priya is writing up her toy count. Write about how you would look at two groups of toys and decide which group is bigger.",
    "11": "Nikhil tipped out his toy box and counted 6 cars, 4 balls and 3 dolls.",
    "11.1": "Altogether, how many toys were in Nikhil's box?",
    "11.2": "Of the three kinds of toy, which did Nikhil own most of, and what was the count?",
    "12": "Priya put her toys into groups and counted 8 blue ones and 5 red ones.",
    "12.1": "By how many do Priya's blue toys outnumber her red ones?",
    "12.2": "Describe the way you found the gap between the two groups.",
    "13": "Read the table below: the toys counted by Nikhil's class were Cars 7, Balls 5, Kites 2.",
    "13.1": "From the table, which toy has the highest count?",
    "13.2": "Using the table, work out the total number of toys counted.",
    "14": "Read this note: Priya's soft toys sit on the shelf while her wooden ones are kept in the box.",
    "14.1": "From the note, in which place are Priya's wooden toys kept?",
    "14.2": "From the note, which rule did Priya follow when sorting her toys?",
}


def rows_for(side):
    out = []
    for sno, typ, (comp, lo), blooms, marks, q, a, e, opts in JA:
        out.append(row(sno, typ, comp, lo, blooms, BLOOM_EXP, marks,
                       PARA[sno] if side == "B" else q,
                       answer=a, explanation=e, options=opts))
    return out


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ja = build_sheet(rows_for("A"), "A", OUTPUT_DIR / "sheet_JA_maths_typologies.xlsx",
                     TEMPLATE, CURRICULUM, {})
    jb = build_sheet(rows_for("B"), "A", OUTPUT_DIR / "sheet_JB_maths_paraphrase.xlsx",
                     TEMPLATE, CURRICULUM, {})
    spec = [{"sno": s[0], "typology": s[1], "A": s[5], "B": PARA[s[0]], "expect": "FLAG"}
            for s in JA]
    (OUTPUT_DIR / "spec_maths_typologies.json").write_text(
        json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
    from collections import Counter
    print(f"JA: {ja}\nJB: {jb}   ({len(JA)} rows)")
    for t, n in sorted(Counter(s[1] for s in JA).items()):
        print(f"   {t:<28} {n}")
    unchanged = [s["sno"] for s in spec if s["A"] == s["B"]]
    print("unchanged questions:", unchanged or "none")


if __name__ == "__main__":
    main()
