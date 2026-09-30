"""Build the Grade 1 Hindi A/C sheets for the QAR duplicate probe.

Mirrors the English pair exactly - same 24-row layout (12 standalone + 4
Case/Source Based parents + 8 sub-rows), same name-swap-only transformation -
so the three subjects are comparable rather than three different experiments.

मीना is NOT in the swap map on purpose: she is the chapter's own character
("CH-1: मीना का परिवार"), so renaming her would change the item's relationship
to its chapter and put Metadata Alignment in play alongside the duplicate
check. Only the reader/child names around her are swapped.
"""
import json
from pathlib import Path

from openpyxl import load_workbook

from tools.name_swap_sheet_builder import build_sheet, make_swapper, row, verify_name_swap

OUTPUT_DIR = Path("data/qar_semantic")
TEMPLATE = Path("data/upload_templates/sme_sheet.xlsx")

CURRICULUM = {
    "grade": "Grade 1",
    "subject": "Hindi",
    "book": "Book 1",
    "unit": None,
    "chapter": "CH-1: मीना का परिवार",
}

NAME_MAP = {
    "रवि": "सुरेश", "अनिता": "प्रिया", "कबीर": "निखिल",
    "सीता": "राधा", "अरुण": "विक्रम", "नेहा": "लता",
}

def _hindi_competencies():
    """Read the Grade 1 Hindi competency/LO pairs out of the template itself.

    Transcribing these as literals does not survive contact with a line break:
    an implicit string concatenation silently inserted one extra space into two
    of them, and the importer compares on exact text, so every row carrying
    those two was rejected. Reading them from the workbook that ships the
    dropdowns removes the class of error entirely.
    """
    workbook = load_workbook(TEMPLATE)
    worksheet = workbook["Items"]
    pairs = []
    for index, excel_row in enumerate(range(2, 21)):
        competency = worksheet.cell(row=excel_row, column=83).value
        if not competency:
            continue
        outcomes = [
            worksheet.cell(row=r, column=92 + index).value
            for r in range(2, worksheet.max_row + 1)
            if worksheet.cell(row=r, column=92 + index).value
        ]
        pairs.append((competency, outcomes[0] if outcomes else None))
    workbook.close()
    return pairs


# Order in the template: 0 write, 1 read aloud, 2 instructions, 3 story
# comprehension, 4 storytelling.
_COMPETENCIES = _hindi_competencies()
C_WRITE, LO_WRITE = _COMPETENCIES[0]
C_READ, LO_READ = _COMPETENCIES[1]
C_INSTR, LO_INSTR = _COMPETENCIES[2]
C_STORY, LO_STORY = _COMPETENCIES[3]
C_TELL, LO_TELL = _COMPETENCIES[4]

ROWS = [
    row("1", "Match the Following", C_STORY, LO_STORY, "Understanding",
        "विद्यार्थी परिवार के शब्दों को उनके अर्थ से जोड़ता है।", "2",
        "रवि परिवार के शब्दों का मिलान कर रहा है। हर शब्द को उसके अर्थ से मिलाइए।",
        answer="1-A, 2-B, 3-C",
        explanation="माता का अर्थ माँ, पिता का अर्थ पापा और दादी का अर्थ पिताजी की माँ है।",
        options=["माता|पिता|दादी", "माँ|पापा|पिताजी की माँ"]),
    row("2", "Match the Following", C_STORY, LO_STORY, "Understanding",
        "विद्यार्थी हर सदस्य को उसके काम से जोड़ता है।", "2",
        "अनिता चित्र कार्ड मिला रही है। हर सदस्य को उसके काम से मिलाइए।",
        answer="1-A, 2-B, 3-C",
        explanation="दादाजी कहानी सुनाते हैं, मीना स्कूल जाती है और माँ खाना बनाती हैं।",
        options=["दादाजी|मीना|माँ", "कहानी सुनाते हैं|स्कूल जाती है|खाना बनाती हैं"]),
    row("3", "True or False", C_STORY, LO_STORY, "Remembering",
        "विद्यार्थी पाठ में दी गई बात को याद करता है।", "1",
        "कबीर ने कहा कि मीना के परिवार में दादा-दादी भी रहते हैं।",
        answer="TRUE",
        explanation="पाठ के अनुसार मीना के घर में दादा-दादी रहते हैं।"),
    row("4", "Fill in the Blank", C_WRITE, LO_WRITE, "Remembering",
        "विद्यार्थी वाक्य पूरा करने के लिए सही शब्द लिखता है।", "1",
        "सीता ने लिखा कि मीना अपने ____ के साथ रहती है।",
        answer="परिवार",
        explanation="मीना अपने परिवार के साथ रहती है, इसलिए रिक्त स्थान में परिवार आएगा।"),
    row("5", "Assertion and Reasoning", C_STORY, LO_STORY, "Analysing",
        "विद्यार्थी जाँचता है कि कारण कथन की व्याख्या करता है या नहीं।", "1",
        "अभिकथन (A): अरुण ने कहा कि मीना का परिवार बड़ा है। कारण (R): उसके परिवार में कई सदस्य रहते हैं।",
        answer="A",
        explanation="दोनों कथन सही हैं और कारण अभिकथन की सही व्याख्या करता है।"),
    row("6", "Assertion and Reasoning", C_STORY, LO_STORY, "Analysing",
        "विद्यार्थी गलत अभिकथन को सही कारण से अलग करता है।", "1",
        "अभिकथन (A): नेहा ने कहा कि मीना अकेली रहती है। कारण (R): परिवार में सब मिलकर रहते हैं।",
        answer="D",
        explanation="मीना अकेली नहीं रहती, इसलिए A गलत है जबकि R सही है।"),
    row("7", "FA Activity", C_INSTR, LO_INSTR, "Applying",
        "विद्यार्थी निर्देशों का पालन करके चित्र बनाता है।", "2",
        "रवि अपने परिवार का चित्र बनाना चाहता है। अपने परिवार का चित्र बनाइए और हर सदस्य का नाम लिखिए।",
        answer="परिवार का चित्र बनाया गया है और हर सदस्य का नाम लिखा गया है।",
        explanation="चित्र बनने और नाम लिखे जाने पर गतिविधि पूरी मानी जाएगी।"),
    row("8", "FA Activity", C_INSTR, LO_INSTR, "Applying",
        "विद्यार्थी दूसरों को स्पष्ट निर्देश देता है।", "2",
        "अनिता कक्षा में खेल कराना चाहती है। परिवार के सदस्यों के नाम बोलकर एक खेल खेलिए और उसके नियम बताइए।",
        answer="खेल खेला गया है और उसके नियम स्पष्ट रूप से बताए गए हैं।",
        explanation="खेल खेलने और नियम बताने पर गतिविधि पूरी मानी जाएगी।"),
    row("9", "Free Response", C_TELL, LO_TELL, "Creating",
        "विद्यार्थी अपने अनुभव से उत्तर रचता है।", "3",
        "कबीर पूछता है कि आपके परिवार में कौन-कौन है। दो-तीन वाक्यों में लिखिए।",
        answer="मेरे परिवार में माँ, पापा, दादी और मैं हैं। हम सब मिलकर रहते हैं।",
        explanation="दो-तीन वाक्यों में दिया गया कोई भी उचित उत्तर स्वीकार्य है।"),
    row("10", "Free Response", C_TELL, LO_TELL, "Evaluating",
        "विद्यार्थी अपनी पसंद बताता है और कारण देता है।", "3",
        "सीता जानना चाहती है कि आपको अपने परिवार के साथ क्या करना अच्छा लगता है। कुछ वाक्यों में लिखिए।",
        answer="मुझे परिवार के साथ खाना खाना और दादी से कहानी सुनना अच्छा लगता है।",
        explanation="कारण सहित दिया गया कोई भी उत्तर स्वीकार्य है।"),
    row("11", "Long Answer Question", C_TELL, LO_TELL, "Understanding",
        "विद्यार्थी पाठ के परिवार का क्रम से वर्णन करता है।", "4",
        "अरुण को कक्षा में मीना के परिवार के बारे में बताना है। मीना के परिवार का वर्णन कीजिए।",
        answer="मीना के परिवार में माँ, पापा, दादा-दादी और छोटा भाई हैं। सब मिलकर रहते हैं और एक-दूसरे की मदद करते हैं।",
        explanation="उत्तर में सदस्यों के नाम और उनका साथ रहना आना चाहिए।"),
    row("12", "Long Answer Question", C_WRITE, LO_WRITE, "Understanding",
        "विद्यार्थी अपने परिवार पर विस्तार से लिखता है।", "4",
        "नेहा अपने परिवार पर लिख रही है। अपने परिवार के बारे में विस्तार से लिखिए।",
        answer="मेरे परिवार में चार सदस्य हैं। माँ खाना बनाती हैं, पापा काम पर जाते हैं और दादी कहानी सुनाती हैं। हम सब मिलकर रहते हैं।",
        explanation="सदस्यों और उनके कामों का वर्णन होने पर उत्तर पूरा माना जाएगा।"),

    # --- Case Based #1 ---
    row("13", "Case Based Question", C_STORY, LO_STORY, "Understanding",
        "विद्यार्थी दिए गए प्रसंग को पढ़कर प्रश्नों के उत्तर देता है।", "3",
        "रवि और उसकी बहन मिलकर पाठ पढ़ रहे थे। वे उस पन्ने पर रुके जहाँ मीना अपने दादा-दादी के साथ बैठी है।",
        explanation="इस प्रसंग पर दो प्रश्न आधारित हैं।"),
    row("13.1", "Multiple Choice Question", C_STORY, LO_STORY, "Remembering",
        "विद्यार्थी प्रसंग से सही जानकारी चुनता है।", "1",
        "रवि के पढ़े पन्ने में मीना किसके साथ बैठी है?",
        answer="दादा-दादी के साथ",
        explanation="प्रसंग के अनुसार मीना दादा-दादी के साथ बैठी है।",
        options=["दादा-दादी के साथ", "मित्रों के साथ", "शिक्षक के साथ", "अकेली"]),
    row("13.2", "Short Answer Question", C_STORY, LO_STORY, "Understanding",
        "विद्यार्थी पात्र के व्यवहार का कारण बताता है।", "2",
        "रवि ने पूछा कि मीना दादा-दादी के पास क्यों बैठी है। कारण लिखिए।",
        answer="वह उनसे कहानी सुनना चाहती है।",
        explanation="दादा-दादी से कहानी सुनना ही उसके पास बैठने का कारण है।"),

    # --- Case Based #2 ---
    row("14", "Case Based Question", C_STORY, LO_STORY, "Understanding",
        "विद्यार्थी दिए गए प्रसंग को पढ़कर प्रश्नों के उत्तर देता है।", "3",
        "अनिता ने विद्यालय में एक नाटक देखा। नाटक में मीना का पूरा परिवार साथ बैठकर खाना खा रहा था।",
        explanation="इस प्रसंग पर दो प्रश्न आधारित हैं।"),
    row("14.1", "True or False", C_STORY, LO_STORY, "Understanding",
        "विद्यार्थी प्रसंग की बात की जाँच करता है।", "1",
        "अनिता के देखे नाटक में मीना का परिवार साथ बैठकर खाना खा रहा था।",
        answer="TRUE",
        explanation="प्रसंग में यही बात कही गई है।"),
    row("14.2", "Very Short Answer Question", C_STORY, LO_STORY, "Remembering",
        "विद्यार्थी प्रसंग की मुख्य क्रिया बताता है।", "2",
        "अनिता के देखे नाटक में परिवार क्या कर रहा था?",
        answer="सब साथ बैठकर खाना खा रहे थे।",
        explanation="प्रसंग के अनुसार पूरा परिवार साथ खाना खा रहा था।"),

    # --- Source Based #1 ---
    row("15", "Source Based Question", C_READ, LO_READ, "Understanding",
        "विद्यार्थी अनुच्छेद पढ़कर प्रश्नों के उत्तर देता है।", "3",
        "सीता ने यह अनुच्छेद पढ़ा: मीना का परिवार बड़ा है। घर में दादा-दादी, माँ, पापा और छोटा भाई रहते हैं। सब मिलकर काम करते हैं और शाम को साथ बैठते हैं।",
        explanation="अनुच्छेद में सदस्यों और उनकी दिनचर्या की जानकारी है।"),
    row("15.1", "Multiple Choice Question", C_READ, LO_READ, "Remembering",
        "विद्यार्थी अनुच्छेद से सही जानकारी चुनता है।", "1",
        "सीता के पढ़े अनुच्छेद में घर में कौन-कौन रहते हैं?",
        answer="दादा-दादी, माँ, पापा और छोटा भाई",
        explanation="अनुच्छेद में यही सदस्य बताए गए हैं।",
        options=["दादा-दादी, माँ, पापा और छोटा भाई", "केवल माँ और पापा",
                 "केवल दादा-दादी", "घर में कोई नहीं रहता"]),
    row("15.2", "Short Answer Question", C_READ, LO_READ, "Understanding",
        "विद्यार्थी अनुच्छेद से दिनचर्या बताता है।", "2",
        "सीता के पढ़े अनुच्छेद के अनुसार परिवार शाम को क्या करता है?",
        answer="सब मिलकर साथ बैठते हैं।",
        explanation="अनुच्छेद में शाम को साथ बैठने की बात कही गई है।"),

    # --- Source Based #2 ---
    row("16", "Source Based Question", C_READ, LO_READ, "Analysing",
        "विद्यार्थी अनुच्छेद में दिए क्रम को समझता है।", "3",
        "कबीर ने ये पंक्तियाँ पढ़ीं: मीना रोज़ सुबह उठकर दादी को नमस्ते करती है। फिर वह तैयार होकर विद्यालय जाती है।",
        explanation="पंक्तियों में मीना की सुबह की दिनचर्या दी गई है।"),
    row("16.1", "Fill in the Blank", C_READ, LO_READ, "Understanding",
        "विद्यार्थी पंक्तियों से सही शब्द भरता है।", "1",
        "कबीर की पढ़ी पंक्तियों में मीना सुबह उठकर दादी को ____ करती है।",
        answer="नमस्ते",
        explanation="पंक्तियों के अनुसार मीना दादी को नमस्ते करती है।"),
    row("16.2", "Very Short Answer Question", C_READ, LO_READ, "Remembering",
        "विद्यार्थी पंक्तियों से अगला कार्य बताता है।", "2",
        "कबीर की पढ़ी पंक्तियों में मीना नमस्ते करने के बाद कहाँ जाती है?",
        answer="विद्यालय",
        explanation="पंक्तियों के अनुसार वह तैयार होकर विद्यालय जाती है।"),
]


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    a = build_sheet(ROWS, "A", OUTPUT_DIR / "sheet_A_hindi.xlsx", TEMPLATE, CURRICULUM, NAME_MAP)
    c = build_sheet(ROWS, "C", OUTPUT_DIR / "sheet_C_hindi_names.xlsx", TEMPLATE, CURRICULUM, NAME_MAP)

    swap = make_swapper(NAME_MAP)
    (OUTPUT_DIR / "hindi_name_pairs.json").write_text(json.dumps(
        {"curriculum": CURRICULUM, "name_map": NAME_MAP,
         "rows": [{**r, "question_A": r["question"], "question_C": swap(r["question"])} for r in ROWS]},
        indent=2, ensure_ascii=False), encoding="utf-8")

    problems = verify_name_swap(ROWS, NAME_MAP)
    print(f"Sheet A: {a}  ({len(ROWS)} rows)")
    print(f"Sheet C: {c}  ({len(ROWS)} rows)")
    print("Name-swap integrity:", "OK - only names differ, every row carries one"
          if not problems else problems)


if __name__ == "__main__":
    main()
