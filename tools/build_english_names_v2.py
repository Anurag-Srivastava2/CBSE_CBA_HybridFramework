"""Post-fix English regression: change ONLY the names, expect duplicates.

Mirrors the pre-fix English pair exactly - 24 rows covering all 12 typologies
twice (12 standalone + 4 Case/Source Based parents + 8 sub-rows, which the
importer folds into 16 items) - so the result is directly comparable to the
pre-fix 2/16.

Fresh content on a chapter the earlier run never touched (CH-5: A Farm) with
names that appear in none of the earlier sheets, because the bank already
holds the CH-4 English items and would otherwise match them instead.

Every one of the 16 items is the SAME QUESTION with a different child's name,
so the correct result is 16/16 flagged as duplicates.
"""
import json
from pathlib import Path

from openpyxl import load_workbook

from tools.name_swap_sheet_builder import build_sheet, make_swapper, row, verify_name_swap

OUTPUT_DIR = Path("data/qar_semantic")
TEMPLATE = Path("data/upload_templates/sme_sheet.xlsx")

CURRICULUM = {
    "grade": "Grade 1", "subject": "English", "book": "Book 1", "unit": None,
    "chapter": "CH-5: A Farm",
}

NAME_MAP = {
    "Nisha": "Kavya", "Farhan": "Rehan", "Deepa": "Shalini",
    "Manoj": "Girish", "Leela": "Sunita", "Aarav": "Dhruv",
}


def _english_curriculum():
    """Competency/LO strings read from the template rather than retyped."""
    wb = load_workbook(TEMPLATE)
    ws = wb["Items"]
    pairs = []
    for index, excel_row in enumerate(range(2, 12)):
        competency = ws.cell(row=excel_row, column=65).value
        if not competency:
            continue
        outcomes = [ws.cell(row=r, column=73 + index).value
                    for r in range(2, ws.max_row + 1)
                    if ws.cell(row=r, column=73 + index).value]
        pairs.append((competency, outcomes))
    wb.close()
    return pairs


_C = _english_curriculum()
STORY, STORY_LO = _C[0][0], _C[0][1][0]
CONV, CONV_LO = _C[1][0], _C[1][1][0]
PHON, PHON_LO = _C[3][0], _C[3][1][0]
VOCAB, VOCAB_LO = _C[4][0], _C[4][1][0]
NARR, NARR_LO = _C[6][0], _C[6][1][0]
INSTR, INSTR_LO = _C[8][0], _C[8][1][0]

ROWS = [
    row("1", "Match the Following", VOCAB, VOCAB_LO, "Understanding",
        "The learner links each farm animal to the sound it makes.", "2",
        "Nisha is matching farm words. Match each animal with the sound it makes.",
        answer="1-A, 2-B, 3-C",
        explanation="A cow moos, a hen clucks and a duck quacks.",
        options=["Cow|Hen|Duck", "Moo|Cluck|Quack"]),
    row("2", "Match the Following", VOCAB, VOCAB_LO, "Understanding",
        "The learner links each farm animal to what it gives us.", "2",
        "Farhan is matching farm pictures. Match each animal with what it gives us.",
        answer="1-A, 2-B, 3-C",
        explanation="A cow gives milk, a hen gives eggs and a sheep gives wool.",
        options=["Cow|Hen|Sheep", "Milk|Eggs|Wool"]),
    row("3", "True or False", STORY, STORY_LO, "Remembering",
        "The learner recalls which animals the farm keeps.", "1",
        "Deepa said that the farmer keeps cows and hens on the farm.",
        answer="TRUE", explanation="The lesson says the farm has cows and hens."),
    row("4", "Fill in the Blank", PHON, PHON_LO, "Remembering",
        "The learner writes the word that completes the sentence.", "1",
        "Manoj wrote that we get milk from a ____.",
        answer="cow", explanation="Milk comes from a cow, so the missing word is cow."),
    row("5", "Assertion and Reasoning", STORY, STORY_LO, "Analysing",
        "The learner judges whether the reason explains the assertion.", "1",
        "Assertion (A): Leela said the hens live on the farm. Reason (R): Hens give us eggs.",
        answer="B",
        explanation="Both statements are true, but giving eggs is not why hens live on the farm."),
    row("6", "Assertion and Reasoning", STORY, STORY_LO, "Analysing",
        "The learner separates a false assertion from a true reason.", "1",
        "Assertion (A): Aarav said the farmer grows wool on the farm. Reason (R): Sheep give us wool.",
        answer="D",
        explanation="Wool is not grown like a crop, so A is false while R is true."),
    row("7", "FA Activity", INSTR, INSTR_LO, "Applying",
        "The learner follows instructions to act out farm animals.", "2",
        "Nisha wants to act out the farm. Make animal masks and show how each animal moves.",
        answer="Masks are made and the movement of each animal is shown.",
        explanation="The activity is complete when the masks are made and the movements shown."),
    row("8", "FA Activity", INSTR, INSTR_LO, "Applying",
        "The learner follows instructions to draw and label a chart.", "2",
        "Farhan is making a farm chart. Draw three farm animals and label each one.",
        answer="Three farm animals are drawn and each one is labelled.",
        explanation="The activity is complete when the drawing carries three labels."),
    row("9", "Free Response", CONV, CONV_LO, "Creating",
        "The learner states a preference and gives a reason.", "3",
        "Deepa asks which farm animal you like best. Write two or three sentences.",
        answer="I like the cow best. It gives us milk and it is very gentle.",
        explanation="Any farm animal is acceptable if the learner gives a reason."),
    row("10", "Free Response", CONV, CONV_LO, "Creating",
        "The learner imagines a day on a farm and describes it.", "3",
        "Manoj wants to know what you would do on a farm for a day. Write a few sentences.",
        answer="I would feed the hens, help milk the cow and play near the field.",
        explanation="Any sensible description of farm activities is acceptable."),
    row("11", "Long Answer Question", NARR, NARR_LO, "Understanding",
        "The learner describes the farm in an ordered way.", "4",
        "Leela has to tell the class about the farm. Describe what you can see on a farm.",
        answer="On a farm you can see cows, hens, ducks and sheep. There are fields of crops, a shed for the animals and a farmer who looks after them all.",
        explanation="The answer should name animals, the fields and the farmer."),
    row("12", "Long Answer Question", NARR, NARR_LO, "Understanding",
        "The learner narrates the farmer's day in order.", "4",
        "Aarav is writing about the farmer's day. Describe what the farmer does from morning to evening.",
        answer="The farmer wakes up early and feeds the animals. Then he goes to the field to work. In the evening he brings the animals back and rests.",
        explanation="The answer should follow the day in order from morning to evening."),

    row("13", "Case Based Question", STORY, STORY_LO, "Understanding",
        "The learner reads a short case and answers questions from it.", "3",
        "Nisha and her brother visited a farm. They stopped near the shed where the cows were being milked.",
        explanation="The case sets up two questions about the milking shed."),
    row("13.1", "Multiple Choice Question", STORY, STORY_LO, "Remembering",
        "The learner recalls what the cows were giving.", "1",
        "In the shed Nisha visited, what were the cows giving?",
        answer="Milk", explanation="Cows are milked, so they were giving milk.",
        options=["Milk", "Eggs", "Wool", "Honey"]),
    row("13.2", "Short Answer Question", STORY, STORY_LO, "Understanding",
        "The learner gives a reason for a daily farm task.", "2",
        "Nisha asked why the farmer milks the cows every morning. Give one reason.",
        answer="So that fresh milk is ready to drink and to sell.",
        explanation="Fresh milk each morning is the reason the lesson gives."),

    row("14", "Case Based Question", STORY, STORY_LO, "Understanding",
        "The learner reads a short case and answers questions from it.", "3",
        "Farhan watched a play about a farm. In one scene the hens ran out of the coop.",
        explanation="The case sets up two questions about the hens."),
    row("14.1", "True or False", STORY, STORY_LO, "Understanding",
        "The learner checks an event against the case.", "1",
        "In the play Farhan watched, the hens ran out of the coop.",
        answer="TRUE", explanation="The case states that the hens ran out of the coop."),
    row("14.2", "Very Short Answer Question", STORY, STORY_LO, "Remembering",
        "The learner names where the hens live.", "2",
        "In the play Farhan watched, where do the hens live?",
        answer="In the coop", explanation="Hens are kept in a coop on the farm."),

    row("15", "Source Based Question", STORY, STORY_LO, "Understanding",
        "The learner reads a passage and answers questions from it.", "3",
        "Deepa read this passage: A farm has many animals. Cows give milk, hens give eggs and sheep give wool. The farmer works hard from morning till evening.",
        explanation="The passage lists the animals and the farmer's working hours."),
    row("15.1", "Multiple Choice Question", STORY, STORY_LO, "Remembering",
        "The learner locates a stated fact in the passage.", "1",
        "In the passage Deepa read, what do sheep give?",
        answer="Wool", explanation="The passage says sheep give wool.",
        options=["Wool", "Milk", "Eggs", "Honey"]),
    row("15.2", "Short Answer Question", STORY, STORY_LO, "Understanding",
        "The learner reports the farmer's working hours.", "2",
        "In the passage Deepa read, how long does the farmer work?",
        answer="From morning till evening.",
        explanation="The passage states the farmer works from morning till evening."),

    row("16", "Source Based Question", STORY, STORY_LO, "Analysing",
        "The learner reads a passage and follows the order of events.", "3",
        "Manoj read these lines: The farmer wakes up early. He feeds the animals and then goes to the field.",
        explanation="The lines give the order of the farmer's morning."),
    row("16.1", "Fill in the Blank", STORY, STORY_LO, "Understanding",
        "The learner completes a sentence from the passage.", "1",
        "In the lines Manoj read, the farmer feeds the ____ before going to the field.",
        answer="animals", explanation="The lines say he feeds the animals first."),
    row("16.2", "Very Short Answer Question", STORY, STORY_LO, "Remembering",
        "The learner names where the farmer goes next.", "2",
        "In the lines Manoj read, where does the farmer go after feeding the animals?",
        answer="To the field", explanation="The lines say he goes to the field."),
]


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    a = build_sheet(ROWS, "A", OUTPUT_DIR / "sheet_A_english_v2.xlsx", TEMPLATE, CURRICULUM, NAME_MAP)
    c = build_sheet(ROWS, "C", OUTPUT_DIR / "sheet_C_english_v2.xlsx", TEMPLATE, CURRICULUM, NAME_MAP)
    swap = make_swapper(NAME_MAP)
    (OUTPUT_DIR / "english_v2_pairs.json").write_text(json.dumps(
        {"curriculum": CURRICULUM, "name_map": NAME_MAP,
         "rows": [{"sno": r["sno"], "typology": r["typology"],
                   "A": r["question"], "C": swap(r["question"])} for r in ROWS]},
        indent=2, ensure_ascii=False), encoding="utf-8")
    problems = verify_name_swap(ROWS, NAME_MAP)
    print(f"A: {a}\nC: {c}  ({len(ROWS)} rows)")
    print("name-swap integrity:",
          "OK - only names differ, every row carries one" if not problems else problems)


if __name__ == "__main__":
    main()
