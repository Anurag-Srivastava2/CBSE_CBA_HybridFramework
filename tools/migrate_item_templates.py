"""Bring the checked-in item-upload templates onto the app's current columns.

The app's template replaced the single `Unit/Theme` column with `Book` and
`Unit`, which pushed `Chapter No._Name` and everything after it one column to
the right. The tester-pack templates under data/ are hand-trimmed copies with a
pre-filled sample row, so they do not pick that up from a fresh download — this
migrates them in place instead of asking anyone to re-cut twelve files by hand.

Idempotent: a sheet that already carries Book and Unit is left alone, so this
can be re-run after the next template change without doubling columns.

    python tools/migrate_item_templates.py [--book "Book 1"] [--dry-run] [path ...]

`sme_sheet.xlsx` keeps the app template's lookup columns and dropdowns, so it
cannot be migrated in place — it is refreshed from a template downloaded out of
the app, with its sample row carried across by header name:

    python tools/migrate_item_templates.py         --from-download ~/Downloads/item-set-template.xlsx         --sample-from data/upload_templates/sme_sheet.xlsx         data/upload_templates/sme_sheet.xlsx
"""
import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.views import BookView

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from utilities.item_template_columns import (  # noqa: E402
    header_columns,
    normalize_header,
    resolve_columns,
)

# Only the hand-trimmed tester-pack templates. sme_sheet.xlsx is deliberately
# not here: it keeps the app template's lookup columns, and openpyxl does not
# move data validations or defined names when a column is inserted, so
# migrating it in place would leave every dropdown pointing one column off.
# Refresh that one from a fresh download instead — see --from-download.
DEFAULT_TARGETS = ("data/typology_templates/*.xlsx",)
DEFAULT_BOOK = "Book 1"


def migrate_worksheet(worksheet, book_value=DEFAULT_BOOK):
    """Insert Book/Unit in place of Unit/Theme. Returns a status string."""
    columns = resolve_columns(worksheet)
    if columns is None:
        return "skipped (not an item-data sheet)"
    if columns["book"] is not None and columns["unit"] is not None:
        return "already migrated"

    headers = header_columns(worksheet)
    legacy_column = headers.get(normalize_header("Unit/Theme"))
    if legacy_column is None:
        return "skipped (no Unit/Theme column to replace)"

    # openpyxl's insert_cols moves cells but not data validations or defined
    # names, so on a sheet carrying either, an in-place insert silently leaves
    # every dropdown and lookup range one column to the left of the data it
    # describes. Refuse rather than produce a workbook that looks right and
    # validates against the wrong columns.
    if worksheet.data_validations.dataValidation:
        return (
            "REFUSED: sheet carries data validations that insert_cols would "
            "misalign — rebuild it from a fresh download instead"
        )

    # Unit/Theme sat exactly where Book sits now, so it is renamed rather than
    # deleted — that keeps the column's width and styling — and Unit is opened
    # up immediately to its right.
    worksheet.cell(row=1, column=legacy_column).value = "Book"
    worksheet.insert_cols(legacy_column + 1)
    worksheet.cell(row=1, column=legacy_column + 1).value = "Unit"
    worksheet.column_dimensions[
        worksheet.cell(row=1, column=legacy_column + 1).column_letter
    ].width = 24

    # Book is mandatory, so it is filled on every row that carries item
    # metadata. Case/Source-Based sub-rows deliberately leave the metadata
    # columns blank and inherit from their parent, and are left blank here too.
    grade_column = columns["grade"]
    filled = 0
    for row in range(2, worksheet.max_row + 1):
        if worksheet.cell(row=row, column=grade_column).value in (None, ""):
            continue
        worksheet.cell(row=row, column=legacy_column).value = book_value
        filled += 1
    return f"migrated (Book filled on {filled} row(s))"


def migrate_workbook(path, book_value=DEFAULT_BOOK, dry_run=False):
    workbook = load_workbook(path)
    results = {}
    if workbook.defined_names:
        return {
            "*": (
                f"REFUSED: {len(workbook.defined_names)} defined name(s) would be "
                "left pointing at the wrong columns — rebuild from a fresh "
                "download instead"
            )
        }
    for worksheet in workbook.worksheets:
        results[worksheet.title] = migrate_worksheet(worksheet, book_value)
    if not dry_run and any(result.startswith("migrated") for result in results.values()):
        workbook.save(path)
    workbook.close()
    return results


def refresh_from_download(download_path, output_path, sample_path=None, book_value=DEFAULT_BOOK):
    """Rebuild a template from a fresh download, keeping its sample row.

    A downloaded template ships with an empty first data row, but most of the
    upload suites copy row 2's metadata down to every row they write, so an
    empty one would upload as a sheet with no grade, subject or chapter. The
    sample row is therefore carried over from `sample_path` *by header name* —
    which is what lets a row written against the pre-Book columns land in the
    right places under the new ones — and Book is filled in, since the old row
    predates it.
    """
    workbook = load_workbook(download_path)
    target_sheet = next(
        (sheet for sheet in workbook.worksheets if resolve_columns(sheet)), None
    )
    if target_sheet is None:
        raise SystemExit(f"{download_path} has no item-data sheet.")
    target_columns = resolve_columns(target_sheet)

    if sample_path:
        sample_workbook = load_workbook(sample_path)
        sample_sheet = next(
            (sheet for sheet in sample_workbook.worksheets if resolve_columns(sheet)), None
        )
        if sample_sheet is None:
            raise SystemExit(f"{sample_path} has no item-data sheet to take a sample from.")
        sample_columns = resolve_columns(sample_sheet)
        for field, source_column in sample_columns.items():
            target_column = target_columns.get(field)
            if source_column is None or target_column is None:
                continue
            target_sheet.cell(row=2, column=target_column).value = sample_sheet.cell(
                row=2, column=source_column
            ).value
        sample_workbook.close()

    if target_columns["book"] is not None:
        target_sheet.cell(row=2, column=target_columns["book"]).value = book_value

    # A downloaded template opens on its Instructions sheet, but every upload
    # fixture in the suite reaches for `workbook.active` and expects the item
    # data. Point the active sheet at Items so the rebuilt template behaves like
    # the single-sheet copy it replaces, without having to drop the Instructions
    # sheet the app ships (and whose version marker the template tests read).
    item_sheet_index = workbook.index(target_sheet)
    workbook.active = item_sheet_index
    # The active tab is persisted through the workbook *view*, not through
    # `workbook.active` alone — and a downloaded template carries no view at
    # all, so without one the setting is dropped on save and the file reopens
    # on Instructions.
    if workbook.views:
        for view in workbook.views:
            view.activeTab = item_sheet_index
    else:
        workbook.views = [BookView(activeTab=item_sheet_index)]

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return target_sheet.title


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="*", help="workbooks to migrate")
    parser.add_argument("--book", default=DEFAULT_BOOK, help="value to write into Book")
    parser.add_argument("--dry-run", action="store_true", help="report without saving")
    parser.add_argument(
        "--from-download",
        help="rebuild the single target from this freshly downloaded template",
    )
    parser.add_argument(
        "--sample-from",
        help="workbook whose row 2 is carried into the rebuilt template",
    )
    args = parser.parse_args(argv)

    if args.from_download:
        if len(args.paths) != 1:
            parser.error("--from-download writes exactly one output path")
        sheet = refresh_from_download(
            args.from_download, args.paths[0], args.sample_from, args.book
        )
        print(f"{Path(args.paths[0]).name:24} {sheet:14} rebuilt from {args.from_download}")
        return 0

    root = Path(__file__).resolve().parent.parent
    if args.paths:
        targets = [Path(path) for path in args.paths]
    else:
        targets = [match for pattern in DEFAULT_TARGETS for match in sorted(root.glob(pattern))]

    if not targets:
        print("No templates matched.")
        return 1
    refused = False
    for target in targets:
        results = migrate_workbook(target, args.book, args.dry_run)
        for sheet, result in results.items():
            print(f"{target.name:24} {sheet:14} {result}")
            refused = refused or result.startswith("REFUSED")
    return 2 if refused else 0


if __name__ == "__main__":
    raise SystemExit(main())
