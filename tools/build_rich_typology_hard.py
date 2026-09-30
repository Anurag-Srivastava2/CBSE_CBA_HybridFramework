"""Rich-typology paraphrase retest, at genuinely hard distance.

The only prior test of these 7 typologies in Maths (build_maths_typology_gap.py,
set IS1360) scored 14/14 - but at 0.55 mean overlap, tier-2 territory that even
the broken pre-fix matcher already handled 4/5 of the time. It proved nothing
about whether rich typologies are semantically matched.

This retest targets ~0.20-0.30 overlap - the band where the simple typologies
(T/F, MCQ, FITB, VSAQ, SAQ) were failing before today's deployment and where
the latest run (IS1407/IS1408) showed real improvement (8/10). If these rich
typologies hold up at the same distance, the fix generalises across typology.
If they collapse, the fix is scoped to the simple typologies only.

Match the Following, Assertion and Reasoning, FA Activity, Free Response and
Long Answer are 5 standalone pairs. Case Based and Source Based are 1 pair
each (parent + 2 sub-questions), so the parent's own duplicate status and its
children's are both exercised - the parent-vs-child gap (BUG-1) is still open
as of the last check on this system.

Chapter CH-11 (Multiplication) is fresh - no earlier sheet used it.
"""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from tools.name_swap_sheet_builder import build_sheet, row

OUTPUT_DIR = Path("data/qar_semantic")
TEMPLATE = Path("data/upload_templates/sme_sheet_r2.xlsx")

CURRICULUM = {
    "grade": "Grade 1", "subject": "Mathematics", "book": "Book 1", "unit": None,
    "chapter": "CH-11: How Many Times? (Multiplication)",
}
BLOOM_EXP = "The learner models equal groups as repeated addition."


def _mult_curriculum():
    wb = load_workbook(TEMPLATE)
    ws = wb["Items"]

    def vals(name):
        ref = wb.defined_names[name].value.split("!")[1].replace("$", "")
        cells = ws[ref]
        if not isinstance(cells, tuple):
            cells = (cells,)
        return [c.value for row_ in cells
                for c in (row_ if isinstance(row_, tuple) else (row_,)) if c.value is not None]

    gi = vals("Grades").index(CURRICULUM["grade"]) + 1
    si = vals(f"G{gi}_Subjects").index(CURRICULUM["subject"]) + 1
    comps = vals(f"G{gi}_S{si}_Competencies")
    ci_mult = next(i for i, c in enumerate(comps, start=1)
                   if c.startswith("Recognises multiplication"))
    ci_prob = next(i for i, c in enumerate(comps, start=1)
                   if c.startswith("Formulates and solves"))
    los_mult = vals(f"G{gi}_S{si}_C{ci_mult}_LOs")
    los_prob = vals(f"G{gi}_S{si}_C{ci_prob}_LOs")
    chapters = vals(f"G{gi}_S{si}_B1_Chapters")
    wb.close()
    if CURRICULUM["chapter"] not in chapters:
        raise SystemExit(f"{CURRICULUM['chapter']!r} not in {chapters}")
    return (comps[ci_mult - 1], los_mult), (comps[ci_prob - 1], los_prob)


MULT, PROB = _mult_curriculum()
LO_GROUPS, LO_PICT, LO_VOCAB = MULT[1][0], MULT[1][1], MULT[1][2]
LO_REAL, LO_SOLVE = PROB[1][0], PROB[1][1]

# typology, (comp, lo), blooms, marks, A, B, answer, explanation, options
PAIRS = [
    ("Match the Following", (MULT[0], LO_VOCAB), "Understanding", "2",
     "Meera is matching groups of fruit to how many are in each group. Match every basket to its count.",
     "For each basket of fruit Meera has, work out how many pieces it holds and match the basket to that number.",
     "1-A, 2-B, 3-C",
     "Basket 1 has 2 apples, basket 2 has 3 apples, basket 3 has 4 apples.",
     ["Basket 1|Basket 2|Basket 3", "2 apples|3 apples|4 apples"]),

    ("Assertion and Reasoning", (MULT[0], LO_GROUPS), "Analysing", "1",
     "Assertion (A): Arjun says 3 groups of 4 pencils is the same as 4 groups of 3 pencils. Reason (R): Both arrangements give the same total count.",
     "Assertion (A): According to Arjun, having 3 groups with 4 pencils each gives the same amount as 4 groups with 3 pencils each. Reason (R): The overall count comes out equal for both.",
     "A", "Both groupings total 12 pencils, and the reason correctly explains why.", None),

    ("FA Activity", (MULT[0], LO_PICT), "Applying", "2",
     "Rani wants to practise multiplication with real objects. Arrange your pencils into equal groups and draw a picture showing what you made.",
     "To get hands-on practice with equal groups, Rani should place her pencils into same-sized groups and sketch what the arrangement looks like.",
     "Pencils are arranged into equal groups and a picture of the grouping is drawn.",
     "The activity is complete once equal groups are formed and pictured.", None),

    ("Free Response", (PROB[0], LO_REAL), "Creating", "3",
     "Kabir asks you to think of a real-life situation where things come in equal groups. Write two or three sentences describing it.",
     "Come up with an everyday example where objects naturally appear in equal-sized groups, and put it into two or three sentences for Kabir.",
     "Eggs come in trays with 6 eggs in each tray, so every tray holds an equal group.",
     "Any real situation involving equal groups is acceptable.", None),

    ("Long Answer Question", (MULT[0], LO_GROUPS), "Understanding", "4",
     "Sana has to explain multiplication to a younger student. Describe how forming equal groups of objects helps in counting a large collection quickly.",
     "A younger learner needs multiplication explained by Sana. Write about the way that splitting a big collection into same-sized groups makes counting it faster.",
     "Splitting objects into equal groups lets you count the number of groups and multiply by the group size instead of counting every object one by one, which is much quicker for a large collection.",
     "The answer should connect equal grouping to faster counting.", None),
]

CABA = {
    "parent_typology": "Case Based Question",
    "comp_lo": (PROB[0], LO_SOLVE),
    "A_parent": "Tara arranged 4 rows of toy cars with 3 cars standing in each row.",
    "B_parent": "In 4 separate rows, Tara lined up toy cars so that each row held 3 cars.",
    "subs": [
        ("Multiple Choice Question", (MULT[0], LO_GROUPS), "Applying", "1",
         "How many toy cars did Tara arrange altogether?",
         "In total, how many toy cars make up Tara's arrangement?",
         "12", "4 rows of 3 cars make 12 cars in all.", ["10", "11", "12", "13"]),
        ("Short Answer Question", (PROB[0], LO_SOLVE), "Analysing", "2",
         "Explain how you worked out the total number of cars.",
         "Describe the method you used to reach the total count of cars.",
         "Multiply the number of rows by the number of cars in each row: 4 times 3 equals 12.",
         "Repeated addition or multiplication of rows by cars-per-row gives the total.", None),
    ],
}

SBQ = {
    "parent_typology": "Source Based Question",
    "comp_lo": (PROB[0], LO_SOLVE),
    "A_parent": "Read the note: A baker places 5 trays of cupcakes on the counter, and every tray carries 6 cupcakes.",
    "B_parent": "Read this: On the counter a baker has set out 5 trays, each one loaded with 6 cupcakes.",
    "subs": [
        ("Very Short Answer Question", (MULT[0], LO_GROUPS), "Applying", "1",
         "According to the note, how many cupcakes are there altogether?",
         "Based on what was read, what is the overall number of cupcakes?",
         "30", "5 trays of 6 cupcakes make 30 cupcakes in all.", None),
        ("Short Answer Question", (PROB[0], LO_SOLVE), "Analysing", "2",
         "According to the note, explain how the total number of cupcakes was found.",
         "From what was read, describe the way the overall cupcake count was worked out.",
         "Multiply the number of trays by the cupcakes on each tray: 5 times 6 equals 30.",
         "Repeated addition or multiplication of trays by cupcakes-per-tray gives the total.", None),
    ],
}


def build_rows(side):
    out = []
    for i, (typ, (comp, lo), blooms, marks, qa, qb, ans, exp, opts) in enumerate(PAIRS, start=1):
        out.append(row(str(i), typ, comp, lo, blooms, BLOOM_EXP, marks,
                       qb if side == "B" else qa, answer=ans, explanation=exp, options=opts))

    next_no = len(PAIRS) + 1
    for block in (CABA, SBQ):
        comp, lo = block["comp_lo"]
        out.append(row(str(next_no), block["parent_typology"], comp, lo, "Understanding", BLOOM_EXP, "3",
                       block["B_parent"] if side == "B" else block["A_parent"],
                       answer=None, explanation="The case/source sets up the sub-questions below."))
        parent_no = next_no
        for j, (typ, (c, l), blooms, marks, qa, qb, ans, exp, opts) in enumerate(block["subs"], start=1):
            out.append(row(f"{parent_no}.{j}", typ, c, l, blooms, BLOOM_EXP, marks,
                           qb if side == "B" else qa, answer=ans, explanation=exp, options=opts))
        next_no = parent_no + 1
    return out


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    qa = build_rows("A")
    qb = build_rows("B")
    ra = build_sheet(qa, "A", OUTPUT_DIR / "sheet_QA_rich.xlsx", TEMPLATE, CURRICULUM, {})
    rb = build_sheet(qb, "A", OUTPUT_DIR / "sheet_QB_rich.xlsx", TEMPLATE, CURRICULUM, {})

    tok = lambda t: set(re.findall(r"[a-z0-9]+", t.lower()))
    spec = []
    for a, b in zip(qa, qb):
        ov = len(tok(a["question"]) & tok(b["question"])) / max(1, len(tok(a["question"]) | tok(b["question"])))
        spec.append({"sno": a["sno"], "typology": a["typology"], "A": a["question"], "B": b["question"],
                     "overlap": round(ov, 2), "is_parent": a["typology"] in
                     ("Case Based Question", "Source Based Question")})
    (OUTPUT_DIR / "spec_rich_typology.json").write_text(
        json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")

    print(f"QA: {ra}\nQB: {rb}   ({len(qa)} rows)\n")
    print(f"{'sno':>6} {'typology':<26} {'ovl':>5}  parent?")
    for s in spec:
        print(f"{s['sno']:>6} {s['typology']:<26} {s['overlap']:>5.2f}  {s['is_parent']}")
    non_parent = [s["overlap"] for s in spec if not s["is_parent"]]
    print(f"\nmean overlap (excl. parents): {sum(non_parent)/len(non_parent):.2f}")


if __name__ == "__main__":
    main()
