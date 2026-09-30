"""Post-deployment retest (today's build): both failing paraphrase types at once.

Two types have been failing in Mathematics and both are re-measured here on
fresh content, in one chapter so they share a single upload:

  numeric      5 pairs - quantities restated as words. Last run IS1362: 0/10.
  number-free  5 pairs - no quantity anywhere.        Last run IS1370: 3/10.

CH-7 (Measurement) hosts both naturally - it carries comparison questions that
need no numbers alongside handspan/cup measurements that do - and no earlier
sheet has used it, so the bank has never seen this content.

Every pair keeps the same answer, so all 10 are the same question reworded and
all 10 must be flagged as duplicates.

Built against the template downloaded today and resolved through its defined
names, because column positions moved when the template went 220 -> 238 wide.
"""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from tools.name_swap_sheet_builder import build_sheet, row

OUTPUT_DIR = Path("data/qar_semantic")
TEMPLATE = Path("data/upload_templates/sme_sheet_today.xlsx")

CURRICULUM = {
    "grade": "Grade 1", "subject": "Mathematics", "book": "Book 1", "unit": None,
    "chapter": "CH-7: Lina’s Family (Measurement)",
}
BLOOM_EXP = "The learner compares and measures lengths, weights and volumes of everyday objects."

NUMBER = (r"\d|\b(one|two|three|four|five|six|seven|eight|nine|ten|eleven|twelve|"
          r"thirteen|fourteen|fifteen|sixteen|seventeen|eighteen|nineteen|twenty|"
          r"thirty|forty|fifty|half|first|second|third)\b")


def _measurement_curriculum():
    wb = load_workbook(TEMPLATE)
    ws = wb["Items"]

    def vals(name):
        ref = wb.defined_names[name].value.split("!")[1].replace("$", "")
        return [c.value for r in ws[ref] for c in r if c.value is not None]

    gi = vals("Grades").index(CURRICULUM["grade"]) + 1
    si = vals(f"G{gi}_Subjects").index(CURRICULUM["subject"]) + 1
    comps = vals(f"G{gi}_S{si}_Competencies")
    ci = next(i for i, c in enumerate(comps, start=1)
              if c.startswith("Performs simple measurements of length"))
    los = vals(f"G{gi}_S{si}_C{ci}_LOs")
    chapters = vals(f"G{gi}_S{si}_B1_Chapters")
    wb.close()
    if CURRICULUM["chapter"] not in chapters:
        raise SystemExit(f"{CURRICULUM['chapter']!r} not in {chapters}")
    return comps[ci - 1], los


COMP, LOS = _measurement_curriculum()
LO_ORDER = next(l for l in LOS if l.startswith("Compares and places objects"))
LO_NEAR = next(l for l in LOS if l.startswith("Distinguishes between near"))
LO_EST = next(l for l in LOS if l.startswith("Estimates short distance"))
LO_VOL = next(l for l in LOS if l.startswith("Estimates and measures volumes"))

# kind, typology, lo, blooms, marks, A, B, answer, explanation, options
PAIRS = [
    ("numeric", "True or False", LO_NEAR, "Understanding", "1",
     "Lina's ribbon measures 9 handspans and her belt measures 4 handspans, so the ribbon is longer.",
     "A ribbon of nine handspans beside a belt of four handspans means the ribbon is the longer item.",
     "TRUE", "9 handspans is more than 4 handspans.", None),
    ("numeric", "Multiple Choice Question", LO_EST, "Applying", "1",
     "A table is 6 handspans wide and a shelf is 8 handspans wide. How much wider is the shelf?",
     "Comparing a shelf of eight handspans with a table of six, by how much does the shelf exceed the table?",
     "2", "8 handspans take away 6 handspans leaves 2.", ["1", "2", "3", "4"]),
    ("numeric", "Fill in the Blank", LO_VOL, "Applying", "1",
     "A jug holds 7 cups of water and a bottle holds 4 cups. The jug holds ____ more cups.",
     "With a jug taking seven cups and a bottle taking four, the extra capacity of the jug is ____ cups.",
     "3", "7 cups take away 4 cups leaves 3 cups.", None),
    ("numeric", "Very Short Answer Question", LO_EST, "Applying", "1",
     "A rope is 12 footsteps long. How many footsteps is half of it?",
     "Half the length of a rope measuring twelve footsteps comes to how many footsteps?",
     "6", "Half of 12 footsteps is 6 footsteps.", None),
    ("numeric", "Short Answer Question", LO_ORDER, "Applying", "2",
     "A stick measures 10 handspans and another measures 6. Explain how you find the difference.",
     "A stick spans ten handspans while another spans six. Describe the way to work out the gap between them.",
     "Take 6 from 10, which leaves a difference of 4 handspans.",
     "Subtracting the shorter measurement from the longer gives the difference.", None),

    ("number-free", "True or False", LO_NEAR, "Understanding", "1",
     "A pencil is longer than a school bus.",
     "The length of a school bus falls short of a pencil's length.",
     "FALSE", "A school bus is far longer than a pencil.", None),
    ("number-free", "Very Short Answer Question", LO_ORDER, "Understanding", "1",
     "Which is heavier, a feather or a stone?",
     "Comparing a feather with a stone, which has the greater weight?",
     "A stone", "A stone weighs more than a feather.", None),
    ("number-free", "Fill in the Blank", LO_VOL, "Understanding", "1",
     "A mug holds less water than a bucket, so the bucket has the larger ____.",
     "Since a bucket takes more water than a mug, the bucket has the bigger ____.",
     "capacity", "The amount a container holds is its capacity.", None),
    ("number-free", "Short Answer Question", LO_EST, "Applying", "2",
     "Explain how you would measure the length of your desk without a ruler.",
     "Describe a way of finding how long your desk is when no ruler is available.",
     "Lay your handspan end to end along the desk and count the handspans.",
     "Repeating a non-standard unit and counting the repeats gives the length.", None),
    ("number-free", "Multiple Choice Question", LO_ORDER, "Applying", "1",
     "Which tool would you use to compare the weight of bags?",
     "To find out which bag weighs more, what tool is needed?",
     "A balance", "A balance compares weights directly.",
     ["A balance", "A ruler", "A cup", "A clock"]),
]


def rows_for(side):
    out = []
    for i, (kind, typ, lo, blooms, marks, qa, qb, ans, exp, opts) in enumerate(PAIRS, start=1):
        out.append(row(str(i), typ, COMP, lo, blooms, BLOOM_EXP, marks,
                       qb if side == "B" else qa,
                       answer=ans, explanation=exp, options=opts))
    return out


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    na = build_sheet(rows_for("A"), "A", OUTPUT_DIR / "sheet_NA_today.xlsx",
                     TEMPLATE, CURRICULUM, {})
    nb = build_sheet(rows_for("B"), "A", OUTPUT_DIR / "sheet_NB_today.xlsx",
                     TEMPLATE, CURRICULUM, {})
    tok = lambda t: set(re.findall(r"[a-z0-9]+", t.lower()))
    spec, wrong = [], []
    for i, (kind, typ, lo, blooms, marks, qa, qb, ans, exp, opts) in enumerate(PAIRS, start=1):
        ov = len(tok(qa) & tok(qb)) / len(tok(qa) | tok(qb))
        free = not (re.search(NUMBER, qa, re.I) or re.search(NUMBER, qb, re.I))
        if (kind == "number-free") != free:
            wrong.append(f"{i} tagged {kind} but number_free={free}")
        spec.append({"sno": str(i), "kind": kind, "typology": typ, "A": qa, "B": qb,
                     "answer": ans, "overlap": round(ov, 2), "number_free": free,
                     "expect": "FLAG"})
    (OUTPUT_DIR / "spec_retest_today.json").write_text(
        json.dumps(spec, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"NA: {na}\nNB: {nb}   ({len(PAIRS)} pairs)\n")
    print(f"{'#':>3} {'kind':<12} {'typology':<26} {'ovl':>5} {'num-free':>9}")
    for s in spec:
        print(f"{s['sno']:>3} {s['kind']:<12} {s['typology']:<26} {s['overlap']:>5.2f} {str(s['number_free']):>9}")
    for k in ("numeric", "number-free"):
        v = [s["overlap"] for s in spec if s["kind"] == k]
        print(f"  {k:<12} mean overlap {sum(v)/len(v):.2f}")
    print("\ntag mismatches:", wrong or "none")


if __name__ == "__main__":
    main()
