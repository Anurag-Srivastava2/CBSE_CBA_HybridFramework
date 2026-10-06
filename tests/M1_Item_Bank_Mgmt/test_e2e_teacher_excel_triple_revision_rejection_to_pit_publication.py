"""E2E: the whole teacher contribution review lifecycle on one item set.

The five uploaded items are split into three lanes so a single upload, a
single QAR cycle and three RWG iterations cover what used to take three
separate uploads:

- Image lane (2 items): sent back at RWG iterations 1 and 2, revised WITH an
  attached image both times, approved at iteration 3, and published.  RWG
  verifies the image before each send-back and every PIT reviewer verifies it
  before voting - which is only observable because this lane survives to PIT.
- Rejection lane (2 items): sent back at iterations 1 and 2, revised without
  an image, then rejected at iteration 3 for exhausting the revision limit.
- Clean lane (1 item): approved at iteration 1 and untouched afterwards.

Both revised lanes also assert the teacher's revised content is visible to the
reviewer after each round, so an item that reaches PIT proves the edit landed
rather than merely that the status moved.
"""
import re
from datetime import datetime
from pathlib import Path
from shutil import copy2
from time import monotonic, sleep
from uuid import uuid4

from openpyxl import load_workbook
import pytest
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait

from pages.common.login_page import LoginPage
from pages.pit.review_queue_page import PITReviewQueuePage
from pages.rwg.review_queue_page import RWGReviewQueuePage
from pages.sme.upload_item_file_page import UploadItemFilePage
from utilities.arithmetic_question_factory import (
    MAX_MIXED_QUESTION_COUNT,
    generate_qar_ready_mixed_questions,
)
from utilities.logger import LogGenerator
from utilities.page_evidence import attach
from utilities.qar_recovery import recover_qar_need_improvement_items
from tests.M1_Item_Bank_Mgmt.m1_surveys import enter_screen, survey_opened_item_set
from utilities.element_checks import ElementChecks
from utilities.item_template_columns import (
    clear_rows_from,
    copy_item_row,
    last_item_column,
    resolve_columns,
    write_row_fields,
)
from utilities.read_config import ReadConfig
from utilities.screenshot_utils import ScreenshotUtils

TEST_IMAGES_FOLDER = Path(__file__).parent.parent.parent / "test_images"


# Item IDs render differently depending on the page: "...-Ch29-i3" right after
# upload (the form every ID in this test is held in), "...-CH-1-i3" on the set
# and review pages. Matching rows on the full upload-time ID found nothing
# there, so a status reader returned {} and RWG iteration 3 "approved" nothing
# - Submit RWG Review then stayed disabled. Rows are matched on the stable
# "IS<number>" set prefix plus the "-iN" item number instead, and handed back
# under the upload-time ID so callers can keep comparing like with like.
_ROW_STATUSES_JS = r"""
const setNumber = arguments[0];
const statusPattern = new RegExp('^(' + arguments[1] + ')$', 'i');
const inRowPattern = new RegExp('\\b(' + arguments[1] + ')\\b', 'i');
const idPattern = new RegExp('IS' + setNumber + '[\\w-]*?-i(\\d+)\\b', 'i');
const result = {};
for (const row of document.querySelectorAll('table tbody tr')) {
    const rowText = (row.innerText || row.textContent || '').trim();
    const id = rowText.match(idPattern);
    if (!id) continue;
    const cells = Array.from(row.querySelectorAll('td'))
        .map(cell => (cell.innerText || cell.textContent || '').trim());
    result[id[1]] = cells.find(value => statusPattern.test(value))
        || (rowText.match(inRowPattern) || [])[1]
        || '';
}
if (Object.keys(result).length) return result;

// Split view: once an item is opened, the set keeps its URL but the items
// render as cards in the left rail instead of a table, so the loop above
// finds nothing and every caller saw {} - no item got actioned and Submit
// RWG Review stayed disabled. Only some cards carry the item ID in their
// text (a revised item shows it, an image-only one shows a placeholder),
// so cards are mapped by their position in the rail: card N is item -iN.
const seenTops = new Set();
const cards = Array.from(document.querySelectorAll('button, a, [role="button"], div'))
    .filter((element) => {
        const rect = element.getBoundingClientRect();
        const text = (element.innerText || '').trim();
        if (!text || rect.width < 80 || rect.height < 35) return false;
        if (rect.left > window.innerWidth * 0.45 || rect.top < 80) return false;
        // The opened item's pane also starts inside the left 45% and its text
        // carries a status badge too, so width is what separates a rail card
        // (narrow) from the detail container (more than half the viewport).
        if (rect.width > window.innerWidth * 0.30) return false;
        // Every rail card carries a marks badge ("1m"). Without this the scan
        // also matched wrappers around the cards and invented positions past
        // the end of the set (a 5-item set reported an i10 and an i11).
        if (!/\b\d+\s*m\b/i.test(text)) return false;
        if (!inRowPattern.test(text)) return false;
        const topKey = Math.round(rect.top / 8) * 8;
        if (seenTops.has(topKey)) return false;
        seenTops.add(topKey);
        return true;
    })
    .sort((left, right) => left.getBoundingClientRect().top - right.getBoundingClientRect().top);
cards.forEach((card, position) => {
    const text = (card.innerText || '').trim();
    const id = text.match(idPattern);
    const status = (text.match(inRowPattern) || [])[1] || '';
    if (status) result[id ? id[1] : String(position + 1)] = status;
});
return result;
"""


def row_statuses(driver, item_set_id, statuses):
    """{"<item_set_id>-iN": status} for the item table on screen."""
    set_number = re.match(r"IS(\d+)", item_set_id, re.IGNORECASE).group(1)
    by_index = driver.execute_script(_ROW_STATUSES_JS, set_number, "|".join(statuses)) or {}
    return {f"{item_set_id}-i{index}": status for index, status in by_index.items()}


def upload_form_item_id(item_set_id, rendered_item_id):
    """A page-rendered item ID ("...-CH-1-i3") in this test's upload-time form."""
    index = re.search(r"-i(\d+)\s*$", rendered_item_id or "", re.IGNORECASE)
    return f"{item_set_id}-i{index.group(1)}" if index else rendered_item_id


def _pick_test_image():
    """Return a real image Path from the shared test_images/ folder."""
    images = sorted(TEST_IMAGES_FOLDER.glob("*.png"))
    assert images, f"No PNG images found in {TEST_IMAGES_FOLDER}"
    return images[0]


class MajorActionEvidenceMixin:
    """Capture submit/reject confirmation dialogs before they are dismissed."""

    def set_evidence_recorder(self, recorder):
        self._evidence_recorder = recorder

    def get_visible_confirmation_summary(self):
        return self.driver.execute_script(
            r"""
            const visible = element => {
                const rect = element.getBoundingClientRect();
                const style = getComputedStyle(element);
                return rect.width > 0 && rect.height > 0
                    && style.display !== 'none' && style.visibility !== 'hidden';
            };
            const selectors = [
                '[role="dialog"]', '[role="alertdialog"]', '[aria-modal="true"]',
                '[class*="modal" i]', '[class*="dialog" i]', '[class*="popup" i]'
            ];
            for (const selector of selectors) {
                for (const element of document.querySelectorAll(selector)) {
                    if (!visible(element)) continue;
                    const text = (element.innerText || element.textContent || '')
                        .replace(/\s+/g, ' ').trim();
                    if (text) return text.slice(0, 220);
                }
            }
            return '';
            """
        )

    def record_confirmation_if_visible(self, action_name):
        recorder = getattr(self, "_evidence_recorder", None)
        if not recorder:
            return ""
        try:
            summary = self.get_visible_confirmation_summary()
        except Exception:
            return ""
        if summary:
            try:
                recorder(f"{action_name} confirmation popup shown: {summary}")
            except Exception:
                pass
        return summary

    def click_required_and_confirm(self, locators, button_name, timeout=15):
        for locator in locators:
            try:
                element = self.wait_utils.until_clickable(locator, timeout=timeout)
                self.driver.execute_script(
                    "arguments[0].scrollIntoView({block: 'center'});",
                    element,
                )
                self.pause_before_action()
                try:
                    element.click()
                except Exception:
                    self.driver.execute_script("arguments[0].click();", element)
                sleep(1)
                self.record_confirmation_if_visible(button_name)
                self.confirm_if_prompted()
                return True
            except Exception:
                continue
        raise TimeoutException(
            f"{button_name} button was not available after all item actions."
        )


class TripleIterationUploadItemFilePage(MajorActionEvidenceMixin, UploadItemFilePage):
    ITEM_INDEX_PATTERN = re.compile(r"-i(\d+)\s*$", re.IGNORECASE)

    @classmethod
    def item_index_of(cls, label):
        """Return the trailing item number of an item label, or None.

        The app names the same item set two different ways - the item header
        renders a chapter-style form (``...-CH-1-i1``) while the timeline and
        item content use the upload form (``...-Ch29-i1``) - and the revision
        loop surfaces whichever the DOM happened to give it.  Only the ``-iN``
        suffix is stable across both, so lane membership and the identities
        reported back to the test are keyed on that rather than on the
        item-set half of the label.
        """
        match = cls.ITEM_INDEX_PATTERN.search(label or "")
        return match.group(1) if match else None

    def revise_items_with_image_lane(self, item_set_id, image_item_ids, image_path):
        """Revise every pending item, attaching image_path only to one lane.

        The revision loop walks whatever the app currently lists as needing
        improvement, so both lanes are edited in the same pass; the callback
        decides per item whether an image goes with the edit.  Returns
        canonical ``{item_set_id}-iN`` labels so the caller can assert on them
        exactly whichever form the app rendered.
        """
        image_indexes = {
            self.item_index_of(item_id)
            for item_id in image_item_ids
            if self.item_index_of(item_id)
        }
        assert len(image_indexes) == len(image_item_ids), (
            f"Could not read an item index from every image-lane ID: {image_item_ids}."
        )

        def edit(label):
            item_index = self.item_index_of(label)
            canonical_label = (
                f"{item_set_id}-i{item_index}" if item_index else label
            )
            self.edit_open_revision_item(
                canonical_label,
                image_path=image_path if item_index in image_indexes else None,
            )
            return canonical_label

        return self._revise_items_loop(item_set_id, edit)

    def submit_uploaded_item_set_for_qar(self, uploaded_file_name=""):
        self._confirmation_action_name = "Submit teacher item set for QAR"
        return super().submit_uploaded_item_set_for_qar(uploaded_file_name)

    def rerun_qar_if_enabled(self):
        self._confirmation_action_name = "Teacher resubmit revised item set"
        return super().rerun_qar_if_enabled()

    def confirm_submit_if_prompted(self):
        try:
            self.wait_utils.until_condition(
                lambda driver: self.get_visible_confirmation_summary() or False,
                timeout=4,
            )
        except TimeoutException:
            pass
        self.record_confirmation_if_visible(
            getattr(self, "_confirmation_action_name", "Submit item set")
        )
        return super().confirm_submit_if_prompted()


class TripleIterationRWGReviewQueuePage(MajorActionEvidenceMixin, RWGReviewQueuePage):
    """Test-local RWG behavior for the terminal third-iteration rejection."""

    REJECT_ITEM_LOCATORS = [
        (By.XPATH, "//button[normalize-space()='Reject' and not(@disabled)]"),
        (
            By.XPATH,
            "//*[self::button or @role='button']"
            "[normalize-space()='Reject' and not(@disabled)]",
        ),
    ]
    REJECT_CONFIRM_LOCATOR = (
        By.XPATH,
        "//*[@role='dialog' or @role='alertdialog' or @aria-modal='true' "
        "or contains(translate(@class,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'modal') "
        "or contains(translate(@class,'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'dialog')]"
        "//*[self::button or @role='button']"
        "[(normalize-space()='Reject' or normalize-space()='Confirm' "
        "or normalize-space()='Yes') and not(@disabled)]",
    )

    def get_reject_button(self):
        def find_open_item_reject(driver):
            return driver.execute_script(
                r"""
                const buttons = Array.from(document.querySelectorAll(
                    'button:not([disabled]), [role="button"]'
                )).filter(button => {
                    const rect = button.getBoundingClientRect();
                    const text = (button.innerText || button.textContent || '').trim();
                    return rect.width > 0
                        && rect.height > 0
                        && rect.left > window.innerWidth * 0.42
                        && /^Reject$/i.test(text)
                        && button.getAttribute('aria-disabled') !== 'true';
                });
                return buttons[0] || null;
                """
            )

        return self.wait_utils.until_condition(find_open_item_reject, timeout=30)

    def confirm_reject_if_prompted(self):
        try:
            confirm_button = self.wait_utils.until_clickable(
                self.REJECT_CONFIRM_LOCATOR,
                timeout=5,
            )
        except TimeoutException:
            return False
        self.driver.execute_script(
            "arguments[0].scrollIntoView({block: 'center'});",
            confirm_button,
        )
        self.record_confirmation_if_visible("Reject item after third RWG iteration")
        try:
            confirm_button.click()
        except Exception:
            self.driver.execute_script("arguments[0].click();", confirm_button)
        return True

    def is_open_item_rejected(self, driver, starting_title):
        if not starting_title:
            return False
        return bool(
            driver.execute_script(
                r"""
                const normalize = value => (value || '')
                    .replace(/\s+/g, ' ')
                    .trim()
                    .toLowerCase();
                const title = normalize(arguments[0]);
                return Array.from(
                    document.querySelectorAll('table tbody tr, button, a, [role="button"], div')
                ).some(element => {
                    const rect = element.getBoundingClientRect();
                    const text = normalize(element.innerText || element.textContent || '');
                    const isTableRow = element.matches('table tbody tr');
                    const isLeftCard = rect.width > 80
                        && rect.width < window.innerWidth * 0.42
                        && rect.height > 30
                        && rect.height < 280
                        && rect.left < window.innerWidth * 0.42;
                    return (isTableRow || isLeftCard)
                        && text.includes(title)
                        && /\brejected\b/i.test(text);
                });
                """,
                starting_title,
            )
        )

    def reject_open_item_after_iteration_limit(self, item_id):
        starting_title = self.get_open_item_title()
        self.mark_all_criteria_no()
        reject_button = self.get_reject_button()
        self.driver.execute_script(
            "arguments[0].scrollIntoView({block: 'center'});",
            reject_button,
        )
        self.pause_before_action()
        try:
            reject_button.click()
        except Exception:
            self.driver.execute_script("arguments[0].click();", reject_button)
        sleep(0.5)
        self.confirm_reject_if_prompted()
        self.wait_utils.until_condition(
            lambda driver: self.is_open_item_rejected(driver, starting_title),
            timeout=30,
        )
        return item_id

    def get_review_item_statuses(self, item_set_id):
        return row_statuses(
            self.driver, item_set_id,
            ("Pending", "Under Review", "Approved", "Revise", "Revised", "Rejected"),
        )

    def reject_revised_and_approve_remaining_as_rwg(
        self,
        item_set_id,
        rejected_item_ids,
        approved_item_ids,
        item_set_url="",
    ):
        self.open_review_item_set(item_set_id, item_set_url)
        # Decide the approve lane from the item table, read now, before any
        # item is opened. Afterwards the page stays in the split view, whose
        # rail is read by card position, and whose badges do not change while
        # the review is open: on QA (2026-10-06, IS1605) i3/i4 still read
        # "Revised" after they were rejected. The table carries each row's own
        # ID and read i1-i4 Revised, i5 Approved: the decision needed.
        statuses_before_actions = self.get_review_item_statuses(item_set_id)
        rejected = []
        for item_id in rejected_item_ids:
            self.return_to_item_set_if_needed(item_set_id)
            self.click_item(item_id)
            rejected.append(self.reject_open_item_after_iteration_limit(item_id))

        if statuses_before_actions:
            row_statuses = statuses_before_actions
        else:
            self.return_to_item_set_if_needed(item_set_id)
            row_statuses = self.get_review_item_statuses(item_set_id)
        pending_approved_ids = [
            item_id
            for item_id in approved_item_ids
            # Iteration 3 reviews items the teacher has just revised, so they
            # sit at "Revised" here - omitting it left the approve lane
            # untouched and Submit RWG Review disabled.
            if row_statuses.get(item_id, "").casefold()
            in {"pending", "under review", "revised"}
        ]
        if pending_approved_ids:
            try:
                self.approve_items_with_yes(item_set_id, pending_approved_ids)
            except TimeoutException as error:
                # Name the decision and the item on screen, so a wrong pick
                # is visible from the error rather than only from a screenshot.
                raise TimeoutException(
                    f"{error.msg} Statuses read {row_statuses}; approving "
                    f"{pending_approved_ids}; open item "
                    f"{self.get_open_item_title()!r}."
                ) from error
        else:
            self.click_item(approved_item_ids[0])

        try:
            self.click_required_and_confirm(
                self.SUBMIT_REVIEW_LOCATORS,
                "Submit RWG Review",
                timeout=30,
            )
        except TimeoutException as error:
            # Submit only enables once every row is actioned, so a disabled
            # button means an item was left alone - say which. The same failure
            # came from a status reader that returned nothing (see the module
            # note on item-ID formats), and that was invisible from the error.
            raise TimeoutException(
                f"{error.msg} Rows now read {self.get_review_item_statuses(item_set_id)}; "
                f"approve lane {approved_item_ids} of which this iteration "
                f"approved {pending_approved_ids}; reject lane {rejected_item_ids}."
            ) from error
        # Return every item RWG leaves approved, not only the ones this
        # iteration had to action: a lane approved in an earlier iteration is
        # already "Approved" here, so it is filtered out of
        # pending_approved_ids while still being PIT-actionable. The caller
        # compares this list against the PIT queue, which shows both.
        return rejected, list(approved_item_ids)


class TripleIterationPITReviewQueuePage(MajorActionEvidenceMixin, PITReviewQueuePage):
    """PIT helper that proves rejected rows never enter the actionable loop."""

    def get_pit_item_statuses(self, item_set_id):
        return row_statuses(
            self.driver, item_set_id,
            ("Pending", "Under Review", "Approved", "Rejected", "Published"),
        )

    def approve_only_expected_items_as_pit(
        self,
        item_set_id,
        expected_approved_item_ids,
        rejected_item_ids,
        item_set_url="",
    ):
        self.open_review_item_set(item_set_id, item_set_url)
        starting_quorum_count = self.get_pit_quorum_approval_count()
        # The PIT page renders IDs as "...-CH-1-iN"; bring them into the
        # upload-time form the expected lists use before comparing.
        all_item_ids = [
            upload_form_item_id(item_set_id, item_id)
            for item_id in self.get_pit_item_ids(item_set_id)
        ]
        pending_item_ids = [
            upload_form_item_id(item_set_id, item_id)
            for item_id in self.get_pending_pit_item_ids(item_set_id)
        ]
        expected = {item_id.casefold() for item_id in expected_approved_item_ids}
        rejected = {item_id.casefold() for item_id in rejected_item_ids}
        actual_pending = {item_id.casefold() for item_id in pending_item_ids}

        assert actual_pending == expected, (
            f"PIT actionable items for {item_set_id} were {pending_item_ids}; "
            f"expected only RWG-approved items {expected_approved_item_ids}."
        )
        assert actual_pending.isdisjoint(rejected), (
            f"Rejected RWG items incorrectly became PIT-actionable: "
            f"{sorted(actual_pending.intersection(rejected))}"
        )

        row_statuses = self.get_pit_item_statuses(item_set_id)
        for rejected_item_id in rejected_item_ids:
            if rejected_item_id in all_item_ids:
                assert row_statuses.get(rejected_item_id, "").casefold() == "rejected", (
                    f"PIT displayed {rejected_item_id} with status "
                    f"{row_statuses.get(rejected_item_id)!r}, not Rejected."
                )

        for item_id in pending_item_ids:
            self.click_item(item_id)
            assert not self.is_pit_item_already_approved(item_id), (
                f"{item_id} was already actioned by this PIT user."
            )
            self.wait_for_pit_decision_panel(item_set_id)
            assert self.has_actionable_pit_vote(), (
                f"Authorise was unavailable for RWG-approved PIT item {item_id}."
            )
            self.authorise_pit_item(item_id)
            self.return_to_item_set_if_needed(item_set_id)

        self.click_item(pending_item_ids[0])
        self.submit_pit_review()
        self.wait_utils.until_condition(
            lambda driver: self.get_pit_quorum_approval_count() > starting_quorum_count
            or self.page_contains_text(driver, "Published"),
            timeout=60,
        )
        return all_item_ids, pending_item_ids, row_statuses

    # What the PIT queue's "Item Set Status" column shows for a set that has
    # some PIT votes but not the quorum. QA shows "Approved" there (IS1606 and
    # IS1607 after vote 1, 2026-10-06), though the queue has no Approved tab;
    # "Pending Review" is what the tab is called. Either means "not yet
    # published", which is the requirement before the last vote.
    AWAITING_QUORUM_STATUSES = ("Approved", "Pending Review")

    def get_item_set_queue_status(self, item_set_id):
        """The set's "Item Set Status" cell in the PIT queue, read by column.

        The row is found by its "IS<number>-" prefix: the queue truncates the
        ID cell ("IS1606-G1-Mathematics-C..."), and it renders the chapter as
        "CH-1" where the upload-time ID says "Ch29", so the full ID never
        matched and this read '' (2026-10-06). The status is then read from
        its own column, since QA shows "Approved" there, a word a pattern
        over the whole row did not list.
        """
        set_prefix = re.match(r"IS\d+", item_set_id, re.IGNORECASE).group(0) + "-"
        return self.driver.execute_script(
            r"""
            const setPrefix = arguments[0].toLowerCase();
            const rowMatches = candidate => (candidate.innerText || candidate.textContent || '')
                .toLowerCase()
                .includes(setPrefix);
            const table = Array.from(document.querySelectorAll('table')).find(
                candidate => Array.from(candidate.querySelectorAll('tbody tr')).some(rowMatches)
            );
            if (!table) return '';
            const row = Array.from(table.querySelectorAll('tbody tr')).find(rowMatches);
            const headers = Array.from(table.querySelectorAll('thead th'))
                .map(th => (th.innerText || th.textContent || '').trim().toLowerCase());
            const column = headers.indexOf('item set status');
            const cells = row.querySelectorAll('td');
            if (column >= 0 && cells[column]) {
                return (cells[column].innerText || cells[column].textContent || '').trim();
            }
            const text = row.innerText || row.textContent || '';
            const match = text.match(/\b(Published|Pending Review|Rejected|Approved)\b/i);
            return match ? match[1] : '';
            """,
            set_prefix,
        )


@pytest.mark.flaky(reruns=1, reruns_delay=5)
# Drives the shared teacher account, the app-assigned RWG reviewer and the same
# three PIT accounts as the other M5 flows, and the portal keeps one active
# session per account - two of these running at once sign each other out
# mid-flow and read each other's queue state.
@pytest.mark.serial
@pytest.mark.usefixtures("setup")
class TestE2ETeacherExcelTripleRevisionRejectionToPITPublication:
    logger = LogGenerator.loggen()
    # Of the MAX_MIXED_QUESTION_COUNT items uploaded, these many are revised
    # with an image and approved at iteration 3, and these many are revised
    # without one and rejected at iteration 3. Whatever is left over is
    # approved at iteration 1 and never touched again.
    IMAGE_LANE_SIZE = 2
    REJECTION_LANE_SIZE = 2

    @classmethod
    def update_teacher_upload_file_questions(cls, upload_file):
        workbook = load_workbook(upload_file)
        worksheet = workbook.active
        columns = resolve_columns(worksheet)
        assert columns, f"{upload_file} has no recognisable item-data sheet."
        max_data_column = last_item_column(worksheet)
        # MAX_MIXED_QUESTION_COUNT (5) pulls in every typology the factory
        # knows about - True or False, Short Answer, Fill in the Blank,
        # Very Short Answer, and MCQ - so this flow proves the triple
        # revision/rejection path is typology-agnostic, not True/False-only.
        mixed_items = generate_qar_ready_mixed_questions(MAX_MIXED_QUESTION_COUNT)
        question_count = len(mixed_items)
        for offset, item in enumerate(mixed_items):
            row = 2 + offset
            if row != 2:
                copy_item_row(worksheet, 2, row, max_data_column)
            write_row_fields(
                worksheet,
                row,
                columns,
                {
                    "sequence": offset + 1,
                    "typology": item["typology"],
                    "question": item["question"],
                    "answer": item["answer"],
                    "explanation": item["explanation"],
                    "marks": item["marks"],
                },
            )
            for option_offset, option_text in enumerate(item["options"]):
                option_column = columns.get(f"option_{option_offset + 1}")
                if option_column is not None:
                    worksheet.cell(row=row, column=option_column).value = option_text
        clear_rows_from(worksheet, 2 + question_count, max_data_column)
        workbook.save(upload_file)
        return question_count

    @classmethod
    def create_unique_teacher_upload_file(cls, source_file_path):
        source_file = Path(source_file_path)
        upload_dir = Path.cwd() / "tmp_uploads"
        upload_dir.mkdir(exist_ok=True)
        unique_file = upload_dir / (
            f"{source_file.stem}_teacher_triple_{uuid4().hex[:12]}{source_file.suffix}"
        )
        copy2(source_file, unique_file)
        return unique_file, cls.update_teacher_upload_file_questions(unique_file)

    def survey_reviewer_screen(self, review_page, role):
        """Survey a reviewer's opened item set, if this test is collecting.

        Read-only: marks no criteria, so it cannot consume the one-time review
        vote the set still needs. Re-published each phase because publish()
        writes the whole accumulated list, so a failure later in this long
        chain still leaves the rows gathered so far on the report card.
        """
        checks = getattr(self, "checks", None)
        if checks is None:
            return
        enter_screen(checks, f"{role} — Opened Item Set")
        survey_opened_item_set(checks, review_page)
        checks.publish()

    def login_as(self, username):
        """Log in with one clean-session retry for a stuck SPA loading shell.

        Catches broadly (not just Selenium's TimeoutException) because a
        stalled WebDriver navigation command surfaces as a raw urllib3
        transport timeout, not a Selenium exception, and would otherwise
        skip this retry entirely.
        """
        page = UploadItemFilePage(self.driver)
        last_error = None
        for attempt in range(2):
            try:
                self.driver.get(ReadConfig.get_base_url())
                LoginPage(self.driver).login_to_application(
                    username,
                    ReadConfig.get_password_for_username(username),
                )
                page.wait_for_application_to_load()
                page.close_popup_if_open()
                return
            except Exception as error:
                last_error = error
                if attempt == 1:
                    break
                try:
                    self.driver.delete_all_cookies()
                    self.driver.execute_script(
                        "window.localStorage.clear(); window.sessionStorage.clear();"
                    )
                except Exception:
                    pass
        raise TimeoutException(
            f"The application did not become ready while logging in as {username}."
        ) from last_error

    # Statuses that mean the set is still inside QAR, so RWG cannot hold it yet.
    QAR_STAGE_MARKERS = ("pending_qar", "pending qar", "qar pending", "qar in progress")

    # Badges meaning "this item has not been revised yet". "Revised" is
    # deliberately absent - it is the done state, and substring-matching it
    # against "Revise" is what would make this read backwards.
    UNREVISED_ITEM_STATUSES = ("revise", "need improvement", "needs improvement",
                               "needs revision", "pending")

    @classmethod
    def unrevised_items(cls, upload_page, item_set_id):
        """Items the set still shows as awaiting the teacher's revision.

        Read from the set's own badges rather than from what the revision loop
        reported: that loop records an item as done as soon as its edit call
        returns, without confirming the badge actually flipped, so its tally
        can be one ahead of the page (observed live: it claimed 4 revised while
        the header read "3 items revised" and one item still badged Revise).
        """
        try:
            statuses = upload_page.get_qar_item_statuses(item_set_id)
        except Exception:
            return {}
        return {
            item_id: status
            for item_id, status in statuses.items()
            if status.strip().casefold() != "revised"
            and any(
                marker in status.casefold() for marker in cls.UNREVISED_ITEM_STATUSES
            )
        }

    @classmethod
    def resubmit_revised_item_set(cls, upload_page, item_set_id, round_number):
        """Resubmit after a teacher revision, and prove the resubmit took.

        rerun_qar_if_enabled() reports a disabled or missing button by
        *returning a string*, not by raising, so an unclicked resubmit used to
        sail straight past this step. The set then sits untouched in the
        teacher's revision bucket and the run failed minutes later at the RWG
        lookup with "not visible in any configured RWG queue" - which reads as
        a product routing defect when in fact nothing was ever sent back.
        Failing here instead names the real cause at the point it happens.
        """
        message = upload_page.rerun_qar_if_enabled()
        normalized = str(message).casefold()
        if "disabled" not in normalized and "not available" not in normalized:
            return message

        # Two very different causes reach here, so say which. An incomplete
        # revision is the app behaving correctly on a set the *test* left
        # half-edited; a missing control with everything revised is the app.
        outstanding = cls.unrevised_items(upload_page, item_set_id)
        if outstanding:
            raise AssertionError(
                f"Teacher revision {round_number} for {item_set_id} could not be "
                f"resubmitted ({message}) because {len(outstanding)} item(s) are "
                f"still awaiting revision: {outstanding}. The revision step "
                "reported them as edited, so the edit did not persist - this is "
                "the revision save, not the resubmit control and not RWG routing."
            )
        raise AssertionError(
            f"Teacher revision {round_number} for {item_set_id} was never "
            f"resubmitted: {message} Every item reads as revised, so the set is "
            "ready to go back and the control itself is the problem. The set "
            "stays in the teacher's revision bucket; this is not RWG routing."
        )

    def wait_for_rwg_handoff(
        self, upload_page, item_set_id, item_set_url, round_number, timeout=420
    ):
        """Wait on the teacher's own view for the resubmitted set to reach RWG.

        A resubmission puts the set back through QAR, and RWG cannot see it
        until that finishes. The RWG queue sweep below only paused 10s between
        passes, so it regularly gave up while the handoff was still in flight
        and reported the set as missing from every queue.

        Returns the detected RWG assignee, or "" when the label never rendered
        or the set left QAR without one. require_assigned_rwg() turns a blank
        into a failure that lists the item statuses; the old fallback of trying
        every RWG account only guessed.
        """
        deadline = monotonic() + timeout
        last_status = ""
        while True:
            # Re-fetched each pass rather than polling the view already on
            # screen: this page does not repaint itself when the backend moves
            # the set on, which is the whole reason the handoff looked absent.
            try:
                upload_page.open_item_set_url_and_wait(item_set_url, item_set_id)
                assignee = upload_page.get_item_set_assignee("rwg")
                if assignee:
                    return assignee
                last_status = (
                    UploadItemFilePage.format_status_summary(
                        upload_page.get_item_set_status_summary()
                    )
                    or last_status
                )
                page_text = self.driver.find_element(By.TAG_NAME, "body").text.casefold()
                in_qar = any(
                    marker in page_text for marker in self.QAR_STAGE_MARKERS
                )
            except Exception as error:
                # Logged, not swallowed: a set page that cannot be read at all
                # otherwise looks identical to one that is simply still in QAR,
                # and this method is deliberately non-fatal, so an unreadable
                # page would leave no trace anywhere.
                self.logger.warning(
                    "Round %s: could not read %s while waiting for the RWG "
                    "handoff: %s",
                    round_number,
                    item_set_id,
                    error,
                )
                in_qar = True
            # Bounded against the deadline, not a flat interval: a check placed
            # before a fixed sleep always overshoots by that interval on the
            # last pass.
            remaining = deadline - monotonic()
            if remaining <= 0:
                self.logger.warning(
                    "Round %s: no RWG assignee was detectable for %s within %ss "
                    "(still showing a QAR stage: %s; statuses: %s).",
                    round_number,
                    item_set_id,
                    timeout,
                    in_qar,
                    last_status or "none read",
                )
                return ""
            if not in_qar:
                # Out of QAR but no assignee shown - the label is simply not
                # rendered, so stop waiting and let the queue sweep find it.
                return ""
            sleep(min(15, remaining))

    def teacher_item_states(self):
        """{item index: (status, typology)} read from the item set table.

        Keyed by the trailing `-iN` index rather than the full ID, because the
        page renders the chapter part of the ID ("CH-1") differently from the
        ID captured at upload ("Ch29").
        """
        rows = self.driver.execute_script(
            r"""
            const statuses = /^(Pending|Under Review|Approved|Revise|Revised|Needs Revision|Rejected|Published)$/i;
            const out = [];
            for (const row of document.querySelectorAll('table tbody tr')) {
                const cells = Array.from(row.querySelectorAll('td'))
                    .map(cell => (cell.innerText || cell.textContent || '').trim());
                if (!cells.length) continue;
                const id = cells[0].replace(/\s+/g, '').match(/-i(\d+)$/i);
                if (!id) continue;
                const status = cells.find(value => statuses.test(value)) || '';
                const typology = cells[cells.indexOf(status) + 1] || '';
                out.push([id[1], status, typology]);
            }
            return out;
            """
        ) or []
        return {index: (status, typology) for index, status, typology in rows}

    # Still waiting on the teacher. After a resubmit that took, the revised
    # items read "Pending" (back in the RWG queue); before one, "Revised".
    UNREVISED_STATUSES = ("revise", "needs revision")

    def assert_revisions_saved(self, item_set_id, expected_item_ids, round_number):
        """No item the teacher was asked to revise may still await revision.

        A revision whose save did not persist leaves that item on "Needs
        Revision", and the set then never goes back to RWG. Checking here
        names the item at the step that broke, instead of failing much later
        with an RWG queue that "never received" the set.
        """
        deadline = monotonic() + 20
        states = self.teacher_item_states()
        while not states and monotonic() < deadline:
            sleep(2)
            states = self.teacher_item_states()
        if not states:
            self.logger.warning(
                "Round %s: could not read item statuses for %s; skipping the "
                "revision-saved check.", round_number, item_set_id,
            )
            return
        unsaved = []
        for item_id in expected_item_ids:
            index = TripleIterationUploadItemFilePage.item_index_of(item_id)
            status, typology = states.get(index, ("not listed", ""))
            if status.casefold() in self.UNREVISED_STATUSES + ("not listed",):
                unsaved.append(f"{item_id} ({typology or 'unknown typology'}): {status}")
        assert not unsaved, (
            f"Teacher revision {round_number} did not save for {len(unsaved)} item(s) - "
            f"{'; '.join(unsaved)}. An item left un-revised keeps {item_set_id} with "
            "the teacher, so it can never return to RWG."
        )

    def require_assigned_rwg(self, upload_page, item_set_id, item_set_url, stage):
        """The RWG named on the teacher's item set page ("Reviewer: rwg N").

        This is the only way the reviewer is chosen - no trying every RWG
        account. A sweep across all of them only guessed, cost extra sign-ins
        against the portal's rate limit, and turned "the set never reached
        RWG" into a misleading "missing from every queue".
        """
        assignee = self.wait_for_rwg_handoff(
            upload_page, item_set_id, item_set_url, round_number=stage
        )
        if not assignee:
            states = self.teacher_item_states()
            raise AssertionError(
                f"{stage}: {item_set_id} shows no \"Reviewer\" on the teacher's item "
                "set page, so it was not handed to any RWG. Item statuses: "
                + (", ".join(f"i{i}={s}" for i, (s, _) in sorted(states.items())) or "unreadable")
            )
        return ReadConfig.get_all_user_username(assignee)

    def open_set_as_assigned_rwg(self, upload_page, item_set_id, rwg_username, item_set_url=""):
        """Sign in once as the assigned RWG and find the set in their queue.

        open_review_item_set() searches the queue for the set first and only
        falls back to its direct URL. The queue can lag the handoff by a few
        seconds, so the lookup is retried - under the same sign-in.
        """
        upload_page.reset_browser_session_to_login()
        self.login_as(rwg_username)
        rwg_queue_page = RWGReviewQueuePage(self.driver)
        last_error = None
        for attempt in range(1, 4):
            try:
                rwg_queue_page.open_review_item_set(item_set_id, item_set_url)
                self.survey_reviewer_screen(rwg_queue_page, "RWG")
                return rwg_username
            except TimeoutException as error:
                last_error = error
                if attempt < 3:
                    sleep(20)
        raise TimeoutException(
            f"The teacher's page names {rwg_username} as reviewer of {item_set_id}, "
            f"but {rwg_username} could not open it from their queue after 3 tries - "
            "the assignment and the reviewer's queue disagree."
        ) from last_error

    def create_fresh_teacher_item_set(
        self,
        request,
        teacher_username,
        upload_page,
        evidence_screenshots,
        max_upload_attempts=3,
    ):
        self.login_as(teacher_username)
        for upload_attempt in range(1, max_upload_attempts + 1):
            upload_file, question_count = self.create_unique_teacher_upload_file(
                ReadConfig.get_upload_item_file_path()
            )
            uploaded_file, upload_message = upload_page.upload_item_file_and_validate(
                upload_file
            )
            self.capture_checkpoint_evidence(
                request,
                evidence_screenshots,
                f"Teacher Excel file validation success message shown: {upload_message}",
                "teacher_file_validation_success_message",
            )
            # The file name lets a wizard that loses the QAR run be resolved
            # from the Sets grid (build #3: QA bounced back to Confirm &
            # Submit with no name to look the set up by).
            item_ids, ocr_message = upload_page.submit_uploaded_item_set_for_qar(
                uploaded_file_name=uploaded_file
            )
            assert len(item_ids) == question_count, (
                f"Fresh upload expected {question_count} IDs but received {len(item_ids)}."
            )
            item_set_id = upload_page.get_item_set_id_from_item_ids(item_ids)
            self.capture_checkpoint_evidence(
                request,
                evidence_screenshots,
                f"QAR completion message shown for {item_set_id}: {ocr_message}",
                "teacher_qar_completion_message",
            )
            try:
                qar_recovery = recover_qar_need_improvement_items(
                    page=upload_page,
                    item_set_id=item_set_id,
                    item_ids=item_ids,
                    workbook_path=upload_file,
                    max_retries=3,
                    run_label="TEACHER-TRIPLE-E2E",
                )
            except AssertionError as error:
                if (
                    "terminal, non-editable failures" not in str(error)
                    or upload_attempt == max_upload_attempts
                ):
                    raise
                self.capture_checkpoint_evidence(
                    request,
                    evidence_screenshots,
                    f"Upload attempt {upload_attempt}/{max_upload_attempts} for "
                    f"{item_set_id} hit a terminal QAR outcome ({error}); "
                    "regenerating a fresh item set and retrying.",
                    f"teacher_qar_terminal_failure_retry_{upload_attempt}",
                )
                continue
            break
        item_set_url = self.driver.current_url
        rwg_username = self.require_assigned_rwg(
            upload_page, item_set_id, item_set_url, "After the first QAR"
        )
        return {
            "item_set_id": item_set_id,
            "item_ids": item_ids,
            "item_set_url": item_set_url,
            "rwg_username": rwg_username,
            "uploaded_file": uploaded_file,
            "question_count": question_count,
            "upload_message": upload_message,
            "ocr_message": ocr_message,
            "initial_qar_recovery": qar_recovery,
            "screenshot": upload_page.capture_sets_verification_screenshot(
                request.node.name
            ),
        }

    @staticmethod
    def assert_exact_item_ids(actual_item_ids, expected_item_ids, label):
        actual = {item_id.casefold() for item_id in actual_item_ids}
        expected = {item_id.casefold() for item_id in expected_item_ids}
        assert actual == expected, (
            f"{label} item IDs were {actual_item_ids}; expected {expected_item_ids}."
        )

    def get_teacher_item_statuses(self, item_set_id):
        return row_statuses(
            self.driver, item_set_id,
            ("Pending", "Under Review", "Approved", "Revise", "Revised", "Rejected", "Published"),
        )

    def get_progress_bar_container(self):
        def find_progress_bar(driver):
            return driver.execute_script(
                r"""
                const visible = element => {
                    const rect = element.getBoundingClientRect();
                    const style = getComputedStyle(element);
                    return rect.width > 0
                        && rect.height > 0
                        && style.display !== 'none'
                        && style.visibility !== 'hidden';
                };
                const exactLabel = (root, label) => Array.from(
                    root.querySelectorAll('*')
                ).some(element => visible(element)
                    && (element.innerText || element.textContent || '').trim() === label);
                const qarLabel = Array.from(document.querySelectorAll('*')).find(
                    element => visible(element)
                        && (element.innerText || element.textContent || '').trim() === 'QAR'
                );
                if (!qarLabel) return null;
                const candidates = [];
                let container = qarLabel.parentElement;
                while (container && container !== document.body) {
                    const rect = container.getBoundingClientRect();
                    if (visible(container)
                        && rect.width >= 180
                        && rect.width <= 600
                        && rect.height >= 45
                        && rect.height <= 220
                        && exactLabel(container, 'QAR')
                        && exactLabel(container, 'RWG')
                        && exactLabel(container, 'PIT')) {
                        candidates.push(container);
                    }
                    container = container.parentElement;
                }
                candidates.sort((left, right) => {
                    const leftRect = left.getBoundingClientRect();
                    const rightRect = right.getBoundingClientRect();
                    return leftRect.width * leftRect.height
                        - rightRect.width * rightRect.height;
                });
                return candidates[0] || null;
                """
            )

        return WebDriverWait(self.driver, 30).until(find_progress_bar)

    def inspect_progress_bar(self):
        progress_bar = self.get_progress_bar_container()
        state = self.driver.execute_script(
            r"""
            const root = arguments[0];
            const visible = element => {
                const rect = element.getBoundingClientRect();
                const style = getComputedStyle(element);
                return rect.width > 0
                    && rect.height > 0
                    && style.display !== 'none'
                    && style.visibility !== 'hidden';
            };
            const parseRgb = value => {
                const match = String(value || '').match(
                    /rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)/i
                );
                return match
                    ? [Number(match[1]), Number(match[2]), Number(match[3])]
                    : null;
            };
            const classify = value => {
                const rgb = parseRgb(value);
                if (!rgb) return '';
                const [red, green, blue] = rgb;
                if (blue >= 180 && blue > green * 1.45 && blue > red * 2) {
                    return 'blue';
                }
                if (green >= 105 && green > red * 1.30 && green > blue * 1.08) {
                    return 'green';
                }
                const spread = Math.max(red, green, blue) - Math.min(red, green, blue);
                if (red >= 90 && red <= 225 && spread <= 45) {
                    return 'gray';
                }
                if (spread > 45) {
                    // Brand-accent (e.g. violet/purple) "in-progress, not yet
                    // confirmed" marker that doesn't fit the blue/green
                    // heuristics above; treated as equivalent to 'green'
                    // by the caller.
                    return 'purple';
                }
                return '';
            };
            const labelCenters = {};
            for (const label of ['QAR', 'RWG', 'PIT']) {
                const element = Array.from(root.querySelectorAll('*')).find(
                    candidate => visible(candidate)
                        && (candidate.innerText || candidate.textContent || '').trim() === label
                );
                if (!element) return {error: 'Missing ' + label + ' label'};
                const rect = element.getBoundingClientRect();
                labelCenters[label] = rect.left + rect.width / 2;
            }

            const candidates = [];
            for (const element of root.querySelectorAll('*')) {
                if (!visible(element)) continue;
                const rect = element.getBoundingClientRect();
                if (rect.width < 12 || rect.width > 42 || rect.height < 4 || rect.height > 20) {
                    continue;
                }
                const style = getComputedStyle(element);
                const values = [
                    style.backgroundColor,
                    style.fill,
                    element.getAttribute('fill'),
                    style.stroke,
                    element.getAttribute('stroke'),
                ];
                for (const value of values) {
                    const color = classify(value);
                    if (!color) continue;
                    candidates.push({
                        color,
                        code: value,
                        x: rect.left + rect.width / 2,
                        key: [
                            Math.round(rect.left),
                            Math.round(rect.top),
                            Math.round(rect.width),
                            Math.round(rect.height),
                            color,
                        ].join(':'),
                    });
                    break;
                }
            }

            const unique = Array.from(
                new Map(candidates.map(candidate => [candidate.key, candidate])).values()
            ).sort((left, right) => left.x - right.x);
            const stages = {QAR: [], RWG: [], PIT: []};
            for (const segment of unique) {
                const stage = Object.keys(labelCenters).sort(
                    (left, right) => Math.abs(segment.x - labelCenters[left])
                        - Math.abs(segment.x - labelCenters[right])
                )[0];
                stages[stage].push({color: segment.color, code: segment.code});
            }
            const checkMarks = root.querySelectorAll(
                'svg path, svg polyline, [data-testid*="check" i], [class*="check" i]'
            ).length;
            return {stages, checkMarks};
            """,
            progress_bar,
        )
        if state.get("error"):
            raise AssertionError(f"Progress bar could not be inspected: {state['error']}.")
        return progress_bar, state

    @staticmethod
    def progress_signature(state):
        return tuple(
            (stage, tuple(segment["color"] for segment in state["stages"][stage]))
            for stage in ("QAR", "RWG", "PIT")
        )

    @staticmethod
    def describe_progress(checkpoint, state):
        stage_details = []
        color_codes = {}
        for stage in ("QAR", "RWG", "PIT"):
            segments = state["stages"][stage]
            wording = ", ".join(
                f"{segment['color']} segment"
                for segment in segments
            )
            stage_details.append(f"{stage} shows {wording}")
            for segment in segments:
                color_codes.setdefault(segment["color"], segment["code"])
        codes = ", ".join(
            f"{color.title()} code {color_codes[color]}"
            for color in ("blue", "green", "gray")
            if color in color_codes
        )
        return f"{checkpoint}: {'; '.join(stage_details)}. {codes}."

    def record_checkpoint_evidence(
        self,
        request,
        evidence_screenshots,
        checkpoint_detail,
        screenshot_path,
    ):
        """File one checkpoint through the shared recorder.

        This used to keep its own list and republish the whole property, which
        dropped every page screenshot `ElementChecks` had filed under the same
        key — including the one taken the moment this test starts. `attach`
        appends to that single list instead, so page visits and checkpoints
        share one numbered series.

        `evidence_screenshots` is still appended to because callers derive the
        `popup_NN` screenshot suffix from its length.
        """
        checkpoint_name = attach(checkpoint_detail, screenshot_path)
        evidence_screenshots.append(
            {"name": checkpoint_name, "path": str(screenshot_path)}
        )
        return str(screenshot_path), checkpoint_name

    def capture_checkpoint_evidence(
        self,
        request,
        evidence_screenshots,
        checkpoint_detail,
        screenshot_suffix,
    ):
        screenshot_path = ScreenshotUtils.capture(
            self.driver,
            f"{request.node.name}_{screenshot_suffix}",
        )
        return self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            checkpoint_detail,
            screenshot_path,
        )

    def record_progress_bar_evidence(
        self,
        request,
        evidence_screenshots,
        checkpoint,
        expected_stage_colors,
        previous_signature=None,
    ):
        progress_bar, state = self.inspect_progress_bar()
        actual_stage_colors = {
            stage: [segment["color"] for segment in state["stages"][stage]]
            for stage in ("QAR", "RWG", "PIT")
        }

        def matches_expected(actual, expected):
            if not expected or len(actual) < len(expected):
                return False

            def satisfies(actual_color, expected_color):
                allowed = (
                    {expected_color}
                    if isinstance(expected_color, str)
                    else set(expected_color)
                )
                if actual_color in allowed:
                    return True
                # The app's "in-progress, not yet confirmed" marker renders
                # as a brand-accent color (classified 'purple') rather than
                # the plain green segment fill; treat it as an alternate
                # rendering of 'green' wherever green is an accepted color.
                return actual_color == "purple" and "green" in allowed

            # How many segments a stage draws follows the workflow's
            # configured iteration count for that stage, which is an
            # environment setting rather than something this flow drives:
            # PIT, for instance, renders two segments where it used to
            # render one. The expectation therefore describes the leading
            # segments this flow actually advances, and any extra trailing
            # segment only has to be a plain not-yet-reached 'gray' or repeat
            # the last expected color.
            trailing = expected[-1]
            trailing_options = (
                {trailing} if isinstance(trailing, str) else set(trailing)
            ) | {"gray"}
            padded = list(expected) + [
                tuple(trailing_options)
            ] * (len(actual) - len(expected))
            return all(
                satisfies(actual_color, expected_color)
                for actual_color, expected_color in zip(actual, padded)
            )

        assert all(
            matches_expected(actual_stage_colors[stage], expected_stage_colors[stage])
            for stage in ("QAR", "RWG", "PIT")
        ), (
            f"{checkpoint} progress colors were {actual_stage_colors}; "
            f"expected segment options {expected_stage_colors}."
        )
        signature = self.progress_signature(state)
        if previous_signature is not None:
            assert signature != previous_signature, (
                f"The progress bar did not change at {checkpoint}: {signature}."
            )

        screenshot_dir = Path.cwd() / "screenshots"
        screenshot_dir.mkdir(exist_ok=True)
        safe_checkpoint = re.sub(r"[^A-Za-z0-9 ]+", "", checkpoint).strip()
        timestamp = datetime.now().strftime("%Y %m %d %H %M %S")
        screenshot_path = screenshot_dir / f"{safe_checkpoint} {timestamp}.png"
        self.driver.execute_script(
            "arguments[0].scrollIntoView({block: 'center'});",
            progress_bar,
        )
        progress_bar.screenshot(str(screenshot_path))
        description = self.describe_progress(checkpoint, state)
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            description,
            screenshot_path,
        )
        return signature, str(screenshot_path), description

    def test_e2e_teacher_third_rwg_iteration_rejects_items_then_pit_publishes_approved_only(
        self,
        request, record_property,
    ):
        # Plain-English orientation for the report, for a reader who does
        # not know this test. One line per step, in the order they happen.
        record_property(
            "test_summary",
            "Sign in as a teacher and upload a fresh Excel workbook holding three "
            "lanes of items: one left clean, one to be revised, one to be rejected.\n"
            "Run three rounds of RWG review, with the teacher editing and "
            "resubmitting between each round.\n"
            "On the third round RWG approves the twice-revised lane and rejects the "
            "other lane for good.\n"
            "Expect PIT to act only on the approved items and publish those, leaving "
            "the rejected ones out of the published set.",
        )
        request.node.user_properties.append(
            (
                "result_checkpoint",
                "fresh teacher Excel upload for the isolated three-lane review lifecycle",
            )
        )
        upload_page = TripleIterationUploadItemFilePage(self.driver)
        rwg_page = TripleIterationRWGReviewQueuePage(self.driver)
        pit_page = TripleIterationPITReviewQueuePage(self.driver)

        # One collector for the whole teacher -> RWG -> PIT chain, re-pointed at
        # each screen as the item set moves through it.
        self.checks = ElementChecks(
            upload_page,
            record_property,
            page_name="Teacher Contribution — Review Lifecycle (image, revision, rejection)",
        )
        self.checks.publish()
        evidence_screenshots = []
        upload_page.set_evidence_recorder(
            lambda detail: self.capture_checkpoint_evidence(
                request,
                evidence_screenshots,
                detail,
                f"popup_{len(evidence_screenshots) + 1:02d}",
            )
        )
        rwg_page.set_evidence_recorder(
            lambda detail: self.capture_checkpoint_evidence(
                request,
                evidence_screenshots,
                detail,
                f"popup_{len(evidence_screenshots) + 1:02d}",
            )
        )
        pit_page.set_evidence_recorder(
            lambda detail: self.capture_checkpoint_evidence(
                request,
                evidence_screenshots,
                detail,
                f"popup_{len(evidence_screenshots) + 1:02d}",
            )
        )
        image_path = _pick_test_image()
        default_teacher = ReadConfig.get_username()
        teacher_username = next(
            (
                username
                for username in ReadConfig.get_role_usernames("teacher")
                if username.casefold() != default_teacher.casefold()
            ),
            default_teacher,
        )

        fresh_set = self.create_fresh_teacher_item_set(
            request,
            teacher_username,
            upload_page,
            evidence_screenshots,
        )
        item_set_id = fresh_set["item_set_id"]
        item_ids = fresh_set["item_ids"]
        item_set_url = fresh_set["item_set_url"]
        rwg_username = fresh_set["rwg_username"]
        evidence_screenshot = fresh_set["screenshot"]
        request.node.user_properties.extend(
            [
                ("item_set_id", item_set_id),
                ("item_ids", ", ".join(item_ids)),
                ("teacher_username", teacher_username),
            ]
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"Fresh item set {item_set_id} is visible with {len(item_ids)} items and moved "
            f"to the first RWG review bucket assigned to {rwg_username}.",
            evidence_screenshot,
        )
        progress_signature, evidence_screenshot, _ = self.record_progress_bar_evidence(
            request,
            evidence_screenshots,
            "Progress bar after QAR and before the first RWG review",
            {
                "QAR": [("blue", "green")],
                "RWG": [("green", "gray"), "gray", "gray"],
                "PIT": ["gray"],
            },
        )

        # Lanes are sliced off the same ordered list that
        # revise_some_items_and_approve_rest slices, so the send-back set at
        # iteration 1 is exactly the image lane plus the rejection lane.
        image_lane_item_ids = item_ids[: self.IMAGE_LANE_SIZE]
        rejection_lane_item_ids = item_ids[
            self.IMAGE_LANE_SIZE : self.IMAGE_LANE_SIZE + self.REJECTION_LANE_SIZE
        ]
        revision_lane_item_ids = image_lane_item_ids + rejection_lane_item_ids
        clean_lane_item_ids = item_ids[len(revision_lane_item_ids) :]
        assert clean_lane_item_ids, (
            f"{item_set_id} produced no lane for RWG to approve at iteration 1: "
            f"{item_ids}."
        )
        request.node.user_properties.extend(
            [
                ("image_lane_item_ids", ", ".join(image_lane_item_ids)),
                ("rejection_lane_item_ids", ", ".join(rejection_lane_item_ids)),
                ("clean_lane_item_ids", ", ".join(clean_lane_item_ids)),
                ("image_path", str(image_path)),
            ]
        )

        request.node.user_properties.append(
            (
                "result_checkpoint",
                "RWG iteration 1 revises the image and rejection lanes and approves the clean lane",
            )
        )
        self.open_set_as_assigned_rwg(
            upload_page, item_set_id, rwg_username, item_set_url
        )
        sent_back_round_1, rwg_approved_item_ids = (
            rwg_page.revise_some_and_approve_rest_as_rwg(
                item_set_id,
                item_ids,
                item_set_url,
                revision_count=len(revision_lane_item_ids),
            )
        )
        self.assert_exact_item_ids(
            sent_back_round_1,
            revision_lane_item_ids,
            "RWG first send-back",
        )
        rejected_item_ids = list(rejection_lane_item_ids)
        pit_allowed_item_ids = list(rwg_approved_item_ids)
        evidence_screenshot = rwg_page.capture_review_screenshot(
            request.node.name,
            "rwg_iteration_1_mixed_result",
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"RWG iteration 1 submitted: {len(sent_back_round_1)} items moved to the "
            f"Teacher revision bucket ({len(image_lane_item_ids)} of them the image lane) "
            f"and {len(rwg_approved_item_ids)} items remained approved.",
            evidence_screenshot,
        )

        request.node.user_properties.append(
            ("result_checkpoint", "teacher edit and resubmit after RWG iteration 1")
        )
        upload_page.reset_browser_session_to_login()
        self.login_as(teacher_username)
        upload_page.open_item_set_url_and_wait(item_set_url, item_set_id)
        progress_signature, evidence_screenshot, _ = self.record_progress_bar_evidence(
            request,
            evidence_screenshots,
            "Progress bar after the first RWG review",
            {
                "QAR": [("blue", "green")],
                "RWG": [
                    ("blue", "green"),
                    ("green", "gray"),
                    "gray",
                ],
                "PIT": ["gray"],
            },
            previous_signature=progress_signature,
        )
        teacher_revised_round_1 = upload_page.revise_items_with_image_lane(
            item_set_id,
            image_lane_item_ids,
            image_path,
        )
        self.assert_exact_item_ids(
            teacher_revised_round_1,
            revision_lane_item_ids,
            "Teacher first revision",
        )
        first_resubmit_message = self.resubmit_revised_item_set(
            upload_page, item_set_id, 1
        )
        upload_page.verify_item_set_from_sets_module(item_set_id, item_ids)
        item_set_url = self.driver.current_url
        evidence_screenshot = upload_page.capture_sets_verification_screenshot(
            f"{request.node.name}_teacher_revision_1_resubmitted"
        )
        self.assert_revisions_saved(item_set_id, revision_lane_item_ids, 1)
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"Teacher revision 1 saved for {len(teacher_revised_round_1)} items, with an "
            f"image attached to the {len(image_lane_item_ids)} image-lane items; "
            f"{first_resubmit_message} The set moved back to the RWG iteration 2 bucket.",
            evidence_screenshot,
        )
        # Wait for the resubmission's own QAR pass to hand the set back to RWG
        # before going looking for it there. Without this the queue sweep below
        # raced the handoff and reported the set missing from every queue.
        rwg_username = self.require_assigned_rwg(
            upload_page, item_set_id, item_set_url, "After teacher revision 1"
        )

        request.node.user_properties.append(
            (
                "result_checkpoint",
                "RWG iteration 2 sees the round-1 edits and image, then sends both lanes back",
            )
        )
        self.open_set_as_assigned_rwg(
            upload_page, item_set_id, rwg_username, item_set_url
        )
        rwg_revised_content_round_1 = rwg_page.verify_revised_items_visible_to_reviewer(
            item_set_id,
            revision_lane_item_ids,
        )
        rwg_image_visibility_round_1 = rwg_page.verify_image_visible_for_items(
            item_set_id,
            image_lane_item_ids,
        )
        request.node.user_properties.extend(
            [
                ("rwg_visible_revised_content_round_1", str(rwg_revised_content_round_1)),
                ("rwg_image_visibility_round_1", str(rwg_image_visibility_round_1)),
            ]
        )
        evidence_screenshot = rwg_page.capture_review_screenshot(
            request.node.name,
            "rwg_iteration_2_revised_content_and_image_visible",
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"RWG opened all {len(revision_lane_item_ids)} revised items before iteration 2: "
            f"the teacher's round-1 text is visible on each and the attached image is "
            f"visible on the {len(image_lane_item_ids)} image-lane items.",
            evidence_screenshot,
        )
        sent_back_round_2 = rwg_page.send_item_set_back_as_rwg(
            item_set_id,
            revision_lane_item_ids,
            item_set_url,
        )
        self.assert_exact_item_ids(
            sent_back_round_2,
            revision_lane_item_ids,
            "RWG second send-back",
        )
        evidence_screenshot = rwg_page.capture_review_screenshot(
            request.node.name,
            "rwg_iteration_2_revision",
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"RWG iteration 2 submitted: the same {len(sent_back_round_2)} items moved "
            "back to the Teacher revision bucket.",
            evidence_screenshot,
        )

        request.node.user_properties.append(
            ("result_checkpoint", "teacher edit and resubmit after RWG iteration 2")
        )
        upload_page.reset_browser_session_to_login()
        self.login_as(teacher_username)
        upload_page.open_item_set_url_and_wait(item_set_url, item_set_id)
        progress_signature, evidence_screenshot, _ = self.record_progress_bar_evidence(
            request,
            evidence_screenshots,
            "Progress bar after the second RWG review",
            {
                "QAR": [("blue", "green")],
                "RWG": [
                    ("blue", "green"),
                    ("blue", "green"),
                    ("green", "gray"),
                ],
                "PIT": ["gray"],
            },
            previous_signature=progress_signature,
        )
        teacher_revised_round_2 = upload_page.revise_items_with_image_lane(
            item_set_id,
            image_lane_item_ids,
            image_path,
        )
        self.assert_exact_item_ids(
            teacher_revised_round_2,
            revision_lane_item_ids,
            "Teacher second revision",
        )
        second_resubmit_message = self.resubmit_revised_item_set(
            upload_page, item_set_id, 2
        )
        upload_page.verify_item_set_from_sets_module(item_set_id, item_ids)
        item_set_url = self.driver.current_url
        evidence_screenshot = upload_page.capture_sets_verification_screenshot(
            f"{request.node.name}_teacher_revision_2_resubmitted"
        )
        self.assert_revisions_saved(item_set_id, revision_lane_item_ids, 2)
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"Teacher revision 2 saved for {len(teacher_revised_round_2)} items, with the "
            f"image re-attached to the {len(image_lane_item_ids)} image-lane items; "
            f"{second_resubmit_message} The set moved back to the final RWG review bucket.",
            evidence_screenshot,
        )
        rwg_username = self.require_assigned_rwg(
            upload_page, item_set_id, item_set_url, "After teacher revision 2"
        )

        request.node.user_properties.append(
            (
                "result_checkpoint",
                "RWG iteration 3 approves the twice-revised image lane and rejects the rejection lane",
            )
        )
        self.open_set_as_assigned_rwg(
            upload_page, item_set_id, rwg_username, item_set_url
        )
        rwg_revised_content_round_2 = rwg_page.verify_revised_items_visible_to_reviewer(
            item_set_id,
            revision_lane_item_ids,
        )
        rwg_image_visibility_round_2 = rwg_page.verify_image_visible_for_items(
            item_set_id,
            image_lane_item_ids,
        )
        request.node.user_properties.extend(
            [
                ("rwg_visible_revised_content_round_2", str(rwg_revised_content_round_2)),
                ("rwg_image_visibility_round_2", str(rwg_image_visibility_round_2)),
            ]
        )
        evidence_screenshot = rwg_page.capture_review_screenshot(
            request.node.name,
            "rwg_iteration_3_revised_content_and_image_visible",
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"RWG opened all {len(revision_lane_item_ids)} twice-revised items before "
            "iteration 3: the teacher's round-2 text is visible on each and the image is "
            f"still visible on the {len(image_lane_item_ids)} image-lane items.",
            evidence_screenshot,
        )
        # Both lanes have been revised twice; iteration 3 decides them
        # differently in one submission - the image lane is approved and
        # carries its attachment on to PIT, the rejection lane is rejected for
        # exhausting the revision limit.
        rwg_rejected_item_ids, final_rwg_approved_item_ids = (
            rwg_page.reject_revised_and_approve_remaining_as_rwg(
                item_set_id,
                rejected_item_ids,
                image_lane_item_ids + pit_allowed_item_ids,
                item_set_url,
            )
        )
        self.assert_exact_item_ids(
            rwg_rejected_item_ids,
            rejected_item_ids,
            "RWG terminal rejection",
        )
        evidence_screenshot = rwg_page.capture_review_screenshot(
            request.node.name,
            "rwg_iteration_3_rejected",
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"RWG iteration 3 submitted: {len(rwg_rejected_item_ids)} repeatedly revised "
            f"items changed to Rejected and {len(final_rwg_approved_item_ids)} approved "
            "items moved to the PIT review bucket.",
            evidence_screenshot,
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"Note: both rejected items ({', '.join(rejected_item_ids)}) are rejected "
            "together by RWG in this single iteration-3 action because they exhausted the "
            "3-iteration revision limit; this is not one item rejected directly by RWG and "
            "a separate item rejected purely for hitting the iteration limit.",
            evidence_screenshot,
        )

        request.node.user_properties.append(
            ("result_checkpoint", "4(a) verify third-iteration teacher items are Rejected by RWG")
        )
        upload_page.reset_browser_session_to_login()
        self.login_as(teacher_username)
        upload_page.open_item_set_url_and_wait(item_set_url, item_set_id)
        progress_signature, evidence_screenshot, _ = self.record_progress_bar_evidence(
            request,
            evidence_screenshots,
            "Progress bar after the third RWG review and rejection decision",
            {
                "QAR": [("blue", "green")],
                "RWG": [
                    ("blue", "green"),
                    ("blue", "green"),
                    ("blue", "green"),
                ],
                "PIT": [("green", "gray")],
            },
            previous_signature=progress_signature,
        )
        status_summary = upload_page.get_item_set_status_summary()
        teacher_statuses = self.get_teacher_item_statuses(item_set_id)
        assert int(status_summary.get("Rejected", 0)) == len(rejected_item_ids), (
            f"Expected {len(rejected_item_ids)} RWG-rejected items after iteration 3; "
            f"status was {upload_page.format_status_summary(status_summary)}."
        )
        for rejected_item_id in rejected_item_ids:
            assert teacher_statuses.get(rejected_item_id, "").casefold() == "rejected", (
                f"{rejected_item_id} was not visibly Rejected after the third RWG iteration: "
                f"{teacher_statuses}."
            )
        request.node.user_properties.append(
            ("rwg_rejected_item_ids", ", ".join(rejected_item_ids))
        )
        evidence_screenshot = upload_page.capture_sets_verification_screenshot(
            f"{request.node.name}_rwg_rejection_verified"
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"Teacher Sets view confirms {len(rejected_item_ids)} items are Rejected; "
            f"item-set status is {upload_page.format_status_summary(status_summary)}.",
            evidence_screenshot,
        )

        completed_pit_approvals = []
        pit_observations = []
        pit_image_visibility = {}
        pit_item_set_status = ""
        # Lane 1, so this suite and the typology E2E can hold quorum at the same
        # time instead of queueing on one worker. See ReadConfig.get_pit_quorum.
        for pit_index, pit_username in enumerate(
            ReadConfig.get_pit_quorum(lane=1), start=1
        ):
            request.node.user_properties.append(
                (
                    "result_checkpoint",
                    f"PIT {pit_index}/3 sees the attached image and acts only on "
                    "RWG-approved items; rejected IDs excluded",
                )
            )
            upload_page.reset_browser_session_to_login()
            self.login_as(pit_username)

            # Read-only pass first: opening items to look at the attachment
            # marks no criteria, so it cannot consume this reviewer's one vote.
            pit_page.open_review_item_set(item_set_id, item_set_url)
            self.survey_reviewer_screen(pit_page, "PIT")
            pit_image_visibility[pit_username] = pit_page.verify_image_visible_for_items(
                item_set_id,
                image_lane_item_ids,
            )
            self.capture_checkpoint_evidence(
                request,
                evidence_screenshots,
                f"PIT reviewer {pit_username} can see the teacher's attached image on all "
                f"{len(image_lane_item_ids)} image-lane items before voting.",
                f"pit_{pit_index}_image_visible",
            )

            all_pit_ids, actionable_pit_ids, pit_statuses = (
                pit_page.approve_only_expected_items_as_pit(
                    item_set_id,
                    final_rwg_approved_item_ids,
                    rejected_item_ids,
                    item_set_url,
                )
            )
            pit_observations.append(
                f"{pit_username}: all={all_pit_ids}, actionable={actionable_pit_ids}, "
                f"statuses={pit_statuses}"
            )
            completed_pit_approvals.append(pit_username)
            pit_item_set_status = pit_page.get_item_set_queue_status(item_set_id)
            expected_pit_statuses = (
                ("Published",) if pit_index == 3
                else TripleIterationPITReviewQueuePage.AWAITING_QUORUM_STATUSES
            )
            assert pit_item_set_status.casefold() in {
                status.casefold() for status in expected_pit_statuses
            }, (
                f"After PIT vote {pit_index}, {item_set_id} status was "
                f"{pit_item_set_status!r}; expected one of {expected_pit_statuses}."
            )
            evidence_screenshot = pit_page.capture_review_screenshot(
                request.node.name,
                f"pit_{pit_index}_approved_only_non_rejected",
            )
            self.record_checkpoint_evidence(
                request,
                evidence_screenshots,
                f"PIT vote {pit_index}/3 submitted by {pit_username}: only "
                f"{len(actionable_pit_ids)} RWG-approved items were actionable; "
                f"rejected items stayed excluded and the set moved to {pit_item_set_status}.",
                evidence_screenshot,
            )

        assert len(completed_pit_approvals) == 3, (
            f"Expected PIT 3/3, completed {completed_pit_approvals}."
        )
        request.node.user_properties.extend(
            [
                ("pit_actionability", " | ".join(pit_observations)),
                ("pit_image_visibility", str(pit_image_visibility)),
            ]
        )
        assert all(
            visible
            for reviewer_results in pit_image_visibility.values()
            for visible in reviewer_results.values()
        ), (
            "Every PIT reviewer must see the teacher's attached image before voting: "
            f"{pit_image_visibility}"
        )

        request.node.user_properties.append(
            ("result_checkpoint", "published set retains rejected items and approved-only PIT path")
        )
        upload_page.reset_browser_session_to_login()
        self.login_as(teacher_username)
        upload_page.open_item_set_url_and_wait(item_set_url, item_set_id)
        final_summary = upload_page.get_item_set_status_summary()
        final_status_text = upload_page.format_status_summary(final_summary)
        progress_signature, evidence_screenshot, final_progress_description = (
            self.record_progress_bar_evidence(
                request,
                evidence_screenshots,
                "Progress bar after PIT publication",
                {
                    "QAR": [("blue", "green")],
                    "RWG": [
                        ("blue", "green"),
                        ("blue", "green"),
                        ("blue", "green"),
                    ],
                    "PIT": [("blue", "green")],
                },
                previous_signature=progress_signature,
            )
        )
        final_state_screenshot = upload_page.capture_sets_verification_screenshot(
            f"{request.node.name}_final_published_status"
        )
        self.record_checkpoint_evidence(
            request,
            evidence_screenshots,
            f"Final Teacher Sets view confirms publication: {final_status_text}; "
            f"{len(rejected_item_ids)} rejected items remained rejected after PIT 3/3.",
            final_state_screenshot,
        )
        evidence_screenshot = final_state_screenshot
        assert pit_item_set_status.casefold() == "published", (
            f"The PIT queue did not show {item_set_id} as Published after vote 3: "
            f"{pit_item_set_status!r}."
        )
        assert int(final_summary.get("Rejected", 0)) == len(rejected_item_ids), (
            f"Rejected count changed after PIT publication: {final_status_text}"
        )
        published_statuses = self.get_teacher_item_statuses(item_set_id)
        for image_item_id in image_lane_item_ids:
            assert published_statuses.get(image_item_id, "").casefold() != "rejected", (
                f"Image-lane item {image_item_id} was expected to survive iteration 3 and "
                f"reach publication, but the Sets view shows: {published_statuses}."
            )
        request.node.user_properties.extend(
            [
                ("post_approval_status", final_status_text),
                (
                    "pit_approval_message",
                    f"Three PIT users approved only the {len(final_rwg_approved_item_ids)} "
                    "non-rejected items.",
                ),
                ("progress_bar_result", final_progress_description),
                ("qar_success_screenshot", evidence_screenshot),
                (
                    "result_description",
                    f"Fresh {item_set_id} passed: the {len(image_lane_item_ids)} image-lane "
                    "items survived two revision rounds with an attached image that RWG and "
                    f"all 3 PIT reviewers could see, RWG rejected {len(rejected_item_ids)} "
                    "teacher items at iteration 3, PIT excluded them, approved only "
                    f"{len(final_rwg_approved_item_ids)} items through 3/3, and published the set.",
                ),
            ]
        )
